import sys
import os
sys.stdout.reconfigure(encoding='utf-8')
import win32com.client
import openpyxl

MASTER_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"
DOWNLOADS_DIR = r"C:\Users\Mohamed Helmy\Downloads"
BACKUP_FILE = r"C:\Users\Mohamed Helmy\Downloads\أرشيف_نسخ_الاحتياط_The_Calm\The_Calm_WMS_V2_BACKUP_BEFORE_PHASE1_CLEANUP.xlsx"

def run_parity_audit():
    print("==================================================")
    print("PHASE 1 RIGOROUS AUDIT & VERIFICATION")
    print("==================================================")

    # 1. Open with openpyxl to check data formulas and values
    print("1. Loading workbooks for mathematical parity...")
    wb_master = openpyxl.load_workbook(MASTER_FILE, data_only=True)
    wb_backup = openpyxl.load_workbook(BACKUP_FILE, data_only=True)

    print(f"   Sheets in Master: {wb_master.sheetnames}")
    expected_sheets = ['أرصدة المخزن', 'الاضافة', 'الصرف', 'الوارد من التشغيل (مستعمل)', 'كارت الصنف', 'أصناف بدون حركة', 'الأذون الملغية']
    for s in expected_sheets:
        assert s in wb_master.sheetnames, f"Missing sheet: {s}"
    print("   -> All required sheets are present!")

    ws_m_bal = wb_master['أرصدة المخزن']
    ws_b_bal = wb_backup['أرصدة المخزن']

    # 2. Check 892 active catalog items and missing units
    print("\n2. Checking catalog items and units...")
    active_items_count = 0
    missing_units = []
    diff_count = 0

    sum_init_m = 0.0
    sum_add_m = 0.0
    sum_dis_m = 0.0
    sum_bal_m = 0.0

    sum_init_b = 0.0
    sum_add_b = 0.0
    sum_dis_b = 0.0
    sum_bal_b = 0.0

    for r in range(4, 8504):
        code_m = ws_m_bal.cell(row=r, column=1).value
        name_m = ws_m_bal.cell(row=r, column=3).value
        unit_m = ws_m_bal.cell(row=r, column=4).value
        init_m = float(ws_m_bal.cell(row=r, column=5).value or 0)
        add_m = float(ws_m_bal.cell(row=r, column=6).value or 0)
        dis_m = float(ws_m_bal.cell(row=r, column=7).value or 0)
        bal_m = float(ws_m_bal.cell(row=r, column=8).value or 0)

        code_b = ws_b_bal.cell(row=r, column=1).value
        init_b = float(ws_b_bal.cell(row=r, column=5).value or 0)
        add_b = float(ws_b_bal.cell(row=r, column=6).value or 0)
        dis_b = float(ws_b_bal.cell(row=r, column=7).value or 0)
        bal_b = float(ws_b_bal.cell(row=r, column=8).value or 0)

        if name_m is not None and str(name_m).strip() != "":
            active_items_count += 1
            if unit_m is None or str(unit_m).strip() == "":
                missing_units.append((r, code_m, name_m))

            sum_init_m += init_m
            sum_add_m += add_m
            sum_dis_m += dis_m
            sum_bal_m += bal_m

            sum_init_b += init_b
            sum_add_b += add_b
            sum_dis_b += dis_b
            sum_bal_b += bal_b

            if bal_m != bal_b or add_m != add_b or dis_m != dis_b or init_m != init_b:
                diff_count += 1
                print(f"   [MISMATCH] Row {r} Code {code_m}: Master Bal={bal_m} vs Backup Bal={bal_b}")

    print(f"   Total Active Catalog Items: {active_items_count}")
    print(f"   Missing Units Count: {len(missing_units)} (Expected: 0)")
    assert len(missing_units) == 0, f"Found missing units: {missing_units}"
    assert active_items_count == 892, f"Active items count mismatch: {active_items_count} != 892"
    assert diff_count == 0, f"Item-by-item balance differences found: {diff_count}"

    print(f"\n3. Grand Parity Check Totals:")
    print(f"   Initial Balance: Master = {sum_init_m:,.2f} | Backup = {sum_init_b:,.2f} | Diff = {sum_init_m - sum_init_b:.2f}")
    print(f"   Total Additions: Master = {sum_add_m:,.2f} | Backup = {sum_add_b:,.2f} | Diff = {sum_add_m - sum_add_b:.2f}")
    print(f"   Total Disbursed: Master = {sum_dis_m:,.2f} | Backup = {sum_dis_b:,.2f} | Diff = {sum_dis_m - sum_dis_b:.2f}")
    print(f"   Current Balance: Master = {sum_bal_m:,.2f} | Backup = {sum_bal_b:,.2f} | Diff = {sum_bal_m - sum_bal_b:.2f}")

    assert abs(sum_bal_m - 15679.0) < 0.001, f"Master balance {sum_bal_m} != 15679.0"
    assert abs(sum_bal_m - sum_bal_b) < 0.001, "Master balance does not match backup balance!"

    # 4. Check Voided Vouchers Count in الاضافة and الصرف
    print("\n4. Checking Voided Vouchers in Movements...")
    ws_add = wb_master['الاضافة']
    code_9999_add = 0
    code_2345_add = 0
    for r in range(7, ws_add.max_row + 1):
        c = ws_add.cell(row=r, column=4).value
        if c == 9999 or str(c).strip() == '9999':
            code_9999_add += 1
        if c == 2345 or str(c).strip() == '2345':
            code_2345_add += 1

    print(f"   الاضافة: Code 9999 count = {code_9999_add} (Expected: 8), Code 2345 count = {code_2345_add} (Expected: 0)")
    assert code_9999_add == 8, f"Expected 8 void vouchers in الاضافة, got {code_9999_add}"
    assert code_2345_add == 0, f"Found unexpected code 2345 in الاضافة: {code_2345_add}"

    ws_dis = wb_master['الصرف']
    code_9999_dis = 0
    code_2345_dis = 0
    for r in range(7, ws_dis.max_row + 1):
        c = ws_dis.cell(row=r, column=5).value
        if c == 9999 or str(c).strip() == '9999':
            code_9999_dis += 1
        if c == 2345 or str(c).strip() == '2345':
            code_2345_dis += 1

    print(f"   الصرف: Code 9999 count = {code_9999_dis} (Expected: 15), Code 2345 count = {code_2345_dis} (Expected: 0)")
    assert code_9999_dis == 15, f"Expected 15 void vouchers in الصرف, got {code_9999_dis}"
    assert code_2345_dis == 0, f"Found unexpected code 2345 in الصرف: {code_2345_dis}"

    # 5. Check Navigation Buttons and Shapes via Excel COM
    print("\n5. Checking Navigation Buttons via Excel COM...")
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb_com = excel.Workbooks.Open(os.path.abspath(MASTER_FILE))
        print("   Excel COM opened master file with 0 errors!")
        s_bal = wb_com.Sheets("أرصدة المخزن")
        s_add = wb_com.Sheets("الاضافة")
        s_dis = wb_com.Sheets("الصرف")
        print(f"   Shapes in 'أرصدة المخزن': {s_bal.Shapes.Count} (Expected: 3)")
        print(f"   Shapes in 'الاضافة': {s_add.Shapes.Count} (Expected: 3)")
        print(f"   Shapes in 'الصرف': {s_dis.Shapes.Count} (Expected: 3)")
        assert s_bal.Shapes.Count >= 3, "Shapes lost in أرصدة المخزن!"
        assert s_add.Shapes.Count >= 3, "Shapes lost in الاضافة!"
        assert s_dis.Shapes.Count >= 3, "Shapes lost in الصرف!"
        wb_com.Close(SaveChanges=False)
    finally:
        excel.Quit()

    # 6. Check Downloads Folder files
    print("\n6. Checking Downloads folder cleanliness...")
    downloads_files = [f for f in os.listdir(DOWNLOADS_DIR) if f.startswith("The_Calm_WMS_V2") and f.endswith(".xlsx")]
    print(f"   Active Master Files matching 'The_Calm_WMS_V2*.xlsx' in Downloads: {downloads_files}")
    assert downloads_files == ["The_Calm_WMS_V2.xlsx"], f"Expected ONLY ['The_Calm_WMS_V2.xlsx'] in Downloads, found: {downloads_files}"

    print("\n==================================================")
    print("ALL AUDIT CHECKS PASSED: 100.00% PARITY & INTEGRITY!")
    print("==================================================")

if __name__ == "__main__":
    run_parity_audit()
