import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

target_path = r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx'
wb = openpyxl.load_workbook(target_path, data_only=True)

error_types = {'#N/A', '#VALUE!', '#REF!', '#NAME?', '#DIV/0!', '#NUM!'}

total_errors = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_errors = []
    r_idx = 0
    for row in ws.iter_rows(values_only=True):
        r_idx += 1
        for c_idx, val in enumerate(row, 1):
            s_val = str(val or '').strip()
            if s_val in error_types:
                sheet_errors.append((r_idx, c_idx, s_val))
    print(f"Sheet '{sheet_name}': {len(sheet_errors)} errors found.")
    if sheet_errors:
        for it in sheet_errors[:5]:
            print(f"   Row {it[0]}, Col {it[1]}: {it[2]}")
    total_errors += len(sheet_errors)

print(f"\nTotal formula evaluation errors across entire workbook: {total_errors}")
if total_errors == 0:
    print(">>> 100% CLEAN WORKBOOK: ZERO ERRORS FOUND! <<<")
