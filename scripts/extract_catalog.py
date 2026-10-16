import sys
import openpyxl
import json

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx', data_only=True)
ws_stock = wb['أرصدة المخزن']

dept_names = [
    (1, "أدوات كهربائية", 4, 503),
    (2, "أدوات سباكة", 504, 1003),
    (3, "أدوات خامات الري", 1004, 1503),
    (4, "حمامات سباحة", 1504, 2003),
    (5, "مستلزمات نظافة", 2004, 2503),
    (6, "أدوات نقاشة", 2504, 3003),
    (7, "أدوات وخامات نجارة", 3004, 3503),
    (8, "أدوات عامة", 3504, 4003),
    (9, "أدوات سيكوريت", 4004, 4503),
    (10, "أدوات الوميتال", 4504, 5003),
    (11, "أدوات مكاتب", 5004, 5503),
    (12, "أدوات ميكانيكا", 5504, 6003),
    (13, "أدوات الأمن وكاميرات المراقبة", 6004, 6503),
    (14, "رخام ومواد بناء", 6504, 7003),
    (15, "ملابس", 7004, 7503),
    (16, "عدد و معدات", 7504, 8003),
    (17, "قسم المستهلكات", 8004, 8503),
]

catalog = []
for dept_id, expected_dept_name, r_start, r_end in dept_names:
    for r in range(r_start, r_end + 1):
        code = ws_stock.cell(r, 1).value
        dept_val = ws_stock.cell(r, 2).value
        name = ws_stock.cell(r, 3).value
        unit = ws_stock.cell(r, 4).value
        open_bal = ws_stock.cell(r, 5).value or 0
        added = ws_stock.cell(r, 6).value or 0
        disb = ws_stock.cell(r, 7).value or 0
        stock = ws_stock.cell(r, 8).value or 0
        
        if name and str(name).strip():
            catalog.append({
                'orig_row': r,
                'orig_dept_id': dept_id,
                'orig_dept_name': str(dept_val).strip() if dept_val else expected_dept_name,
                'orig_code': str(code).strip() if code is not None else '',
                'name': str(name).strip(),
                'unit': str(unit).strip() if unit else '',
                'open_bal': float(open_bal),
                'added': float(added),
                'disb': float(disb),
                'stock': float(stock)
            })

print(f"Total active named items found: {len(catalog)}")
with open(r'scripts\scratch\extracted_raw_catalog.json', 'w', encoding='utf-8') as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)

print("Saved raw catalog to scripts/scratch/extracted_raw_catalog.json")
