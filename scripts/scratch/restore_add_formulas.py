# -*- coding: utf-8 -*-
import openpyxl
import win32com.client
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'

print("Loading workbook to restore formulas in الاضافة...")
wb = openpyxl.load_workbook(excel_path)
ws_add = wb['الاضافة']

restored_add_e = 0
restored_add_f = 0

for r in range(7, ws_add.max_row + 1):
    code_val = ws_add.cell(r, 4).value
    if code_val is None or str(code_val).strip() == '':
        continue
    if str(code_val).strip() == '9999': # Skip canceled vouchers
        continue
        
    e_val = str(ws_add.cell(r, 5).value or '')
    f_val = str(ws_add.cell(r, 6).value or '')
    
    if not e_val.startswith('='):
        ws_add.cell(r, 5).value = f'=IFERROR(VLOOKUP(D{r},\'أرصدة المخزن\'!$A:$E,3,0),"")'
        restored_add_e += 1
        
    if not f_val.startswith('='):
        ws_add.cell(r, 6).value = f'=IFERROR(VLOOKUP(D{r},\'أرصدة المخزن\'!$A:$H,4,0),"")'
        restored_add_f += 1

print(f"Restored in الاضافة:")
print(f"  Col E (اسم الصنف): {restored_add_e} rows restored to dynamic VLOOKUP")
print(f"  Col F (الوحدة): {restored_add_f} rows restored to dynamic VLOOKUP")

wb.save(excel_path)
wb.close()

print("Saved. Running Excel COM recalculation...")
excel_app = win32com.client.DispatchEx('Excel.Application')
excel_app.Visible = False
excel_app.DisplayAlerts = False
try:
    wb_com = excel_app.Workbooks.Open(excel_path)
    wb_com.Save()
    wb_com.Close(SaveChanges=True)
    print("Recalculation and COM save complete!")
finally:
    excel_app.Quit()
