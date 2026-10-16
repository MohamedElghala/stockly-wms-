import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws = wb['أرصدة المخزن']

all_items = []
unit_distribution = {}
unit_counts = {}

for r in range(4, ws.max_row+1):
    c_code = ws.cell(row=r, column=1).value
    c_dept = ws.cell(row=r, column=2).value
    c_name = ws.cell(row=r, column=3).value
    c_unit = str(ws.cell(row=r, column=4).value or 'عدد').strip()
    c_init = float(ws.cell(row=r, column=5).value or 0)
    c_in = float(ws.cell(row=r, column=6).value or 0)
    c_out = float(ws.cell(row=r, column=7).value or 0)
    c_bal = float(ws.cell(row=r, column=8).value or 0)
    
    if c_name and str(c_name).strip():
        it = {
            'row': r,
            'code': str(c_code),
            'dept': str(c_dept),
            'name': str(c_name).strip(),
            'unit': c_unit,
            'init': c_init,
            'in': c_in,
            'out': c_out,
            'bal': c_bal
        }
        all_items.append(it)
        unit_distribution[c_unit] = unit_distribution.get(c_unit, 0) + c_bal
        unit_counts[c_unit] = unit_counts.get(c_unit, 0) + 1

print(f"Total Active Items: {len(all_items)}")
total_raw_bal = sum(it['bal'] for it in all_items)
print(f"Total Stock Raw Balance: {total_raw_bal:,.2f}")

print("\n=== Current Balance Breakdown by Unit ===")
for u, tot in sorted(unit_distribution.items(), key=lambda x: x[1], reverse=True):
    cnt = unit_counts[u]
    pct = (tot / total_raw_bal) * 100
    print(f"  الوحدة: {u:<10} | الأصناف: {cnt:>3} | إجمالي الرصيد: {tot:>10,.2f} ({pct:>5.1f}%)")

print("\n=== All Items Measured in Meters (متر) and their Packaging Rules ===")

# Categorize meter items
pipes_komar = []
pipes_castle = []
pipes_poly = []
pipes_pe = []
wires = []
hoses_irrigation = []
other_meters = []

for it in all_items:
    if 'متر' in it['unit']:
        name = it['name']
        dept = it['dept']
        
        # Komar Pipe (6m)
        if 'ماسورة' in name and ('كومر' in name or 'بوصة' in name) and 'كاسل' not in name and 'بولي' not in name:
            pipes_komar.append(it)
        # Castle Pipe (3m)
        elif 'ماسورة' in name and 'كاسل' in name:
            pipes_castle.append(it)
        # Poly PPR Pipe (4m)
        elif 'ماسورة' in name and ('بولي' in name or 'PPR' in name or 'ppr' in name) and 'ايثيلين' not in name and 'إيثيلين' not in name:
            pipes_poly.append(it)
        # PE Pipe (6m)
        elif 'ماسورة' in name and ('ايثيلين' in name or 'إيثيلين' in name or 'PE' in name):
            pipes_pe.append(it)
        # Wires & Cables (100m)
        elif 'سلك' in name or 'كابل' in name:
            wires.append(it)
        # Irrigation Hoses
        elif 'خرطوم' in name and 'ري' in dept:
            hoses_irrigation.append(it)
        else:
            other_meters.append(it)

print(f"1. مواسير كومر (طول العود 6م): {len(pipes_komar)} أصناف")
for it in pipes_komar:
    m = it['bal']
    full = int(m // 6)
    rem = m % 6
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر -> {full} عود مكتمل (+ {rem} م كسر)")

print(f"\n2. مواسير كاسل (طول العود 3م): {len(pipes_castle)} أصناف")
for it in pipes_castle:
    m = it['bal']
    full = int(m // 3)
    rem = m % 3
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر -> {full} عود مكتمل (+ {rem} م كسر)")

print(f"\n3. مواسير بولي (طول العود 4م): {len(pipes_poly)} أصناف")
for it in pipes_poly:
    m = it['bal']
    full = int(m // 4)
    rem = m % 4
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر -> {full} عود مكتمل (+ {rem} م كسر)")

print(f"\n4. مواسير بولي إيثيلين (طول العود 6م): {len(pipes_pe)} أصناف")
for it in pipes_pe:
    m = it['bal']
    full = int(m // 6)
    rem = m % 6
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر -> {full} عود مكتمل (+ {rem} م كسر)")

print(f"\n5. أسلاك وكابلات كهربائية (اللفة 100م): {len(wires)} أصناف")
for it in wires:
    m = it['bal']
    full = int(m // 100)
    rem = m % 100
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر -> {full} لفة مكتملة (+ {rem} م كسر)")

print(f"\n6. خراطيم ري سادة وخراطيم زراعة: {len(hoses_irrigation)} أصناف")
for it in hoses_irrigation:
    m = it['bal']
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {m} متر")

print(f"\n7. أمتار أخرى: {len(other_meters)} أصناف")
for it in other_meters:
    print(f"   - كود {it['code']} | {it['name']} | رصيد: {it['bal']} {it['unit']}")
