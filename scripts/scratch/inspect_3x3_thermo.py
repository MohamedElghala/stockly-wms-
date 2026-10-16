# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws_stock = wb['أرصدة المخزن']

print("--- فحص الأصناف المشابهة لسلك 3*3 في أرصدة المخزن ---")
for r in range(4, 504):
    c = ws_stock.cell(r, 1).value
    n = str(ws_stock.cell(r, 3).value or '')
    u = str(ws_stock.cell(r, 4).value or '')
    op = ws_stock.cell(r, 5).value or 0
    in_v = ws_stock.cell(r, 6).value or 0
    out_v = ws_stock.cell(r, 7).value or 0
    bal = ws_stock.cell(r, 8).value or 0
    
    if '3*3' in n or '3 * 3' in n or 'ترمو' in n or ('سلك' in n and '3' in n):
        print(f"صف {r:3d} | كود: {c} | اسم: {n:<35} | وحدة: {u:<6} | وارد: {in_v:<6} | صرف: {out_v:<6} | رصيد: {bal}")
