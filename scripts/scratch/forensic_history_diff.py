import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BACKUP_DIR = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"
CURRENT_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"

files = [
    ("V2.1_3_SEP09", os.path.join(BACKUP_DIR, "The_Calm_WMS_V2.1 (3).xlsx")),
    ("XLSB_SEP13", os.path.join(BACKUP_DIR, "The_Calm_WMS_V2 (version 1).xlsb.xlsx")),
    ("BEFORE_709", os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx")),
    ("BEFORE_RESTRUCT", os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx")),
    ("AFTER_RESTRUCT", os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx")),
    ("CURRENT", CURRENT_FILE),
]

def safe_float(val):
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val or 0).strip())
    except Exception:
        return -999999.0

def load_tx(label, path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws_c = wb["أرصدة المخزن"]
    ws_a = wb["الاضافة"]
    ws_d = wb["الصرف"]

    cat = {}
    for r_idx, r in enumerate(ws_c.iter_rows(min_row=4, max_row=8503, values_only=True), 4):
        code, dept, name, unit, op, in_q, out_q, bal = (list(r) + [None]*8)[:8]
        if name and str(name).strip():
            cat[r_idx] = (code, str(dept or "").strip(), str(name).strip(), str(unit or "").strip(), safe_float(op), safe_float(in_q), safe_float(out_q), safe_float(bal))

    adds = {}
    for r_idx, r in enumerate(ws_a.iter_rows(min_row=7, max_row=2000, values_only=True), 7):
        seq, dt, v, code, name, unit, qty, t, note = (list(r) + [None]*9)[:9]
        if code is not None or qty is not None or v is not None:
            q_val = safe_float(qty)
            if q_val == -999999.0:
                print(f"  [WEIRD QTY in {label} Add Row {r_idx}]: {r[:9]}")
            adds[r_idx] = (seq, str(dt)[:10] if dt else None, v, code, str(name or "").strip(), unit, q_val, note)

    disbs = {}
    for r_idx, r in enumerate(ws_d.iter_rows(min_row=7, max_row=3000, values_only=True), 7):
        seq, dt, ent, v, code, name, unit, qty, t, note, bal = (list(r) + [None]*11)[:11]
        if code is not None or qty is not None or v is not None:
            q_val = safe_float(qty)
            if q_val == -999999.0:
                print(f"  [WEIRD QTY in {label} Disb Row {r_idx}]: {r[:11]}")
            disbs[r_idx] = (seq, str(dt)[:10] if dt else None, str(ent or "").strip(), v, code, str(name or "").strip(), unit, q_val, note)

    wb.close()
    return cat, adds, disbs

snaps = {}
for label, path in files:
    snaps[label] = load_tx(label, path)
    print(f"Loaded {label}: {len(snaps[label][0])} cat, {len(snaps[label][1])} adds, {len(snaps[label][2])} disbs")

def compare_snaps(l1, l2):
    print(f"\n==========================================")
    print(f"COMPARING {l1} -> {l2}")
    c1, a1, d1 = snaps[l1]
    c2, a2, d2 = snaps[l2]

    print(f"--- Additions Diff ({l1} vs {l2}) ---")
    add_diffs = 0
    for r in sorted(set(a1.keys()) & set(a2.keys())):
        seq1, dt1, v1, cd1, n1, u1, q1, note1 = a1[r]
        seq2, dt2, v2, cd2, n2, u2, q2, note2 = a2[r]
        if n1 != n2 or abs(q1 - q2) > 1e-6 or v1 != v2:
            add_diffs += 1
            print(f"  Add Row {r:4d} (V={v1}->{v2}): [{cd1}] '{n1}' (qty={q1})  -->  [{cd2}] '{n2}' (qty={q2})")
    print(f"Total Addition Row Changes: {add_diffs}")

    print(f"--- Disbursements Diff ({l1} vs {l2}) ---")
    dis_diffs = 0
    for r in sorted(set(d1.keys()) & set(d2.keys())):
        seq1, dt1, ent1, v1, cd1, n1, u1, q1, note1 = d1[r]
        seq2, dt2, ent2, v2, cd2, n2, u2, q2, note2 = d2[r]
        if n1 != n2 or abs(q1 - q2) > 1e-6 or v1 != v2:
            dis_diffs += 1
            print(f"  Disb Row {r:4d} (V={v1}->{v2}): [{cd1}] '{n1}' (qty={q1})  -->  [{cd2}] '{n2}' (qty={q2})")
    print(f"Total Disbursement Row Changes: {dis_diffs}")

compare_snaps("V2.1_3_SEP09", "XLSB_SEP13")
compare_snaps("XLSB_SEP13", "BEFORE_709")
compare_snaps("BEFORE_709", "BEFORE_RESTRUCT")
compare_snaps("BEFORE_RESTRUCT", "AFTER_RESTRUCT")
compare_snaps("AFTER_RESTRUCT", "CURRENT")
