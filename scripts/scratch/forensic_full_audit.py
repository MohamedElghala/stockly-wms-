import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BACKUP_DIR = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"
CURRENT_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"

BEFORE_RESTRUCTURE = os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx")
XLSB_ORIG = os.path.join(BACKUP_DIR, "The_Calm_WMS_V2 (version 1).xlsb.xlsx")
BEFORE_709 = os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx")
SINGLE_BACKUP = os.path.join(BACKUP_DIR, "The_Calm_WMS_V2_SINGLE_SAFETY_BACKUP.xlsx")

print("=== 1. CHECKING FORMULAS & DIRECT SUMS IN CURRENT WORKBOOK ===")
wb_val = openpyxl.load_workbook(CURRENT_FILE, data_only=True, read_only=True)
wb_frm = openpyxl.load_workbook(CURRENT_FILE, data_only=False, read_only=True)

ws_cat_v = wb_val["أرصدة المخزن"]
ws_cat_f = wb_frm["أرصدة المخزن"]
ws_add_v = wb_val["الاضافة"]
ws_add_f = wb_frm["الاضافة"]
ws_dis_v = wb_val["الصرف"]
ws_dis_f = wb_frm["الصرف"]

# Load catalog items
cat_by_code = {}
cat_by_name = {}
formula_issues_cat = []

for r_idx, (rv, rf) in enumerate(zip(
    ws_cat_v.iter_rows(min_row=4, max_row=8503, values_only=True),
    ws_cat_f.iter_rows(min_row=4, max_row=8503, values_only=True)
), 4):
    code_v, dept_v, name_v, unit_v, op_v, in_v, out_v, bal_v = (list(rv) + [None]*8)[:8]
    code_f, dept_f, name_f, unit_f, op_f, in_f, out_f, bal_f = (list(rf) + [None]*8)[:8]
    if name_v and str(name_v).strip():
        c_int = int(code_v) if isinstance(code_v, (int, float)) else str(code_v).strip()
        n_str = str(name_v)
        cat_by_code[c_int] = {
            "row": r_idx, "code": c_int, "dept": dept_v, "name": n_str, "unit": unit_v,
            "open": float(op_v or 0), "in": float(in_v or 0), "out": float(out_v or 0), "bal": float(bal_v or 0)
        }
        if n_str.strip() in cat_by_name:
            print(f"  [DUPLICATE NAME IN CATALOG] '{n_str}' at row {r_idx} and {cat_by_name[n_str.strip()]['row']}")
        cat_by_name[n_str.strip()] = cat_by_code[c_int]

        # Check formulas in F, G, H
        if not str(in_f or "").startswith("=SUMIFS"):
            formula_issues_cat.append((r_idx, c_int, n_str, "IN_FORMULA", in_f))
        if not str(out_f or "").startswith("=SUMIFS"):
            formula_issues_cat.append((r_idx, c_int, n_str, "OUT_FORMULA", out_f))
        if not str(bal_f or "").startswith("="):
            formula_issues_cat.append((r_idx, c_int, n_str, "BAL_FORMULA", bal_f))

print(f"Catalog active items: {len(cat_by_code)} | Formula issues in Catalog: {len(formula_issues_cat)}")
for issue in formula_issues_cat[:20]:
    print("  Cat Formula Issue:", issue)

# Check all rows in الاضافة
add_sum_by_code = {}
add_issues = []
for r_idx, (rv, rf) in enumerate(zip(
    ws_add_v.iter_rows(min_row=7, max_row=2000, values_only=True),
    ws_add_f.iter_rows(min_row=7, max_row=2000, values_only=True)
), 7):
    seq_v, dt_v, v_v, code_v, name_v, unit_v, qty_v, type_v, note_v = (list(rv) + [None]*9)[:9]
    seq_f, dt_f, v_f, code_f, name_f, unit_f, qty_f, type_f, note_f = (list(rf) + [None]*9)[:9]
    if code_v is not None or qty_v is not None or v_v is not None:
        c_int = int(code_v) if isinstance(code_v, (int, float)) else str(code_v or "").strip()
        if c_int in (9999, 2345, "9999", "2345", ""):
            continue
        q_val = float(qty_v or 0)
        add_sum_by_code[c_int] = add_sum_by_code.get(c_int, 0.0) + q_val
        if not str(name_f or "").startswith("="):
            add_issues.append((r_idx, v_v, c_int, name_v, "HARDCODED_NAME_IN_ADD", name_f))
        if "اضاف" not in str(type_v or "") and "إضاف" not in str(type_v or ""):
            add_issues.append((r_idx, v_v, c_int, name_v, "MISSING_TYPE_IN_ADD", type_v))
        if c_int not in cat_by_code:
            add_issues.append((r_idx, v_v, c_int, name_v, "CODE_NOT_IN_CATALOG", q_val))
        else:
            cat_name = cat_by_code[c_int]["name"]
            if str(name_v or "") != cat_name:
                add_issues.append((r_idx, v_v, c_int, name_v, f"NAME_MISMATCH_VS_CAT({cat_name})", q_val))

print(f"Addition issues found: {len(add_issues)}")
for issue in add_issues[:30]:
    print("  Add Issue:", issue)

# Check all rows in الصرف
dis_sum_by_code = {}
dis_issues = []
for r_idx, (rv, rf) in enumerate(zip(
    ws_dis_v.iter_rows(min_row=7, max_row=3000, values_only=True),
    ws_dis_f.iter_rows(min_row=7, max_row=3000, values_only=True)
), 7):
    seq_v, dt_v, ent_v, v_v, code_v, name_v, unit_v, qty_v, type_v, note_v, bal_v = (list(rv) + [None]*11)[:11]
    seq_f, dt_f, ent_f, v_f, code_f, name_f, unit_f, qty_f, type_f, note_f, bal_f = (list(rf) + [None]*11)[:11]
    if code_v is not None or qty_v is not None or v_v is not None:
        c_int = int(code_v) if isinstance(code_v, (int, float)) else str(code_v or "").strip()
        if c_int in (9999, 2345, "9999", "2345", ""):
            continue
        q_val = float(qty_v or 0)
        dis_sum_by_code[c_int] = dis_sum_by_code.get(c_int, 0.0) + q_val
        if not str(name_f or "").startswith("="):
            dis_issues.append((r_idx, v_v, c_int, name_v, "HARDCODED_NAME_IN_DIS", name_f))
        if c_int not in cat_by_code:
            dis_issues.append((r_idx, v_v, c_int, name_v, "CODE_NOT_IN_CATALOG", q_val))
        else:
            cat_name = cat_by_code[c_int]["name"]
            if str(name_v or "") != cat_name:
                dis_issues.append((r_idx, v_v, c_int, name_v, f"NAME_MISMATCH_VS_CAT({cat_name})", q_val))

print(f"Disbursement issues found: {len(dis_issues)}")
for issue in dis_issues[:30]:
    print("  Disb Issue:", issue)

# Compare Python direct sum by code vs Excel SUMIFS value in أرصدة المخزن
sum_mismatches = []
for c_int, info in cat_by_code.items():
    py_in = add_sum_by_code.get(c_int, 0.0)
    py_out = dis_sum_by_code.get(c_int, 0.0)
    py_bal = info["open"] + py_in - py_out
    if abs(py_in - info["in"]) > 1e-6 or abs(py_out - info["out"]) > 1e-6 or abs(py_bal - info["bal"]) > 1e-6:
        sum_mismatches.append((info["row"], c_int, info["name"], info["in"], py_in, info["out"], py_out, info["bal"], py_bal))

print(f"Direct Code Sum vs Excel SUMIFS mismatches: {len(sum_mismatches)}")
for m in sum_mismatches:
    print("  SUM MISMATCH:", m)

wb_val.close()
wb_frm.close()
