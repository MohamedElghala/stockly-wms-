import openpyxl
import json
import os

EXCEL_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"
OUTPUT_JSON = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\wms-monorepo\apps\web\public\data\live_items_snapshot.json"

wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb['أرصدة المخزن']

items = []
sum_bal = 0.0

for r in range(4, ws.max_row + 1):
    code = ws.cell(row=r, column=1).value
    dept = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    unit = ws.cell(row=r, column=4).value
    init_bal = ws.cell(row=r, column=5).value
    added = ws.cell(row=r, column=6).value
    disb = ws.cell(row=r, column=7).value
    bal = ws.cell(row=r, column=8).value

    if name and str(name).strip() and dept and str(dept).strip():
        code_str = str(int(code) if isinstance(code, (int, float)) else str(code).strip())
        name_str = str(name).strip()
        dept_str = str(dept).strip()
        unit_str = str(unit).strip() if unit else 'عدد'
        bal_num = float(bal) if bal is not None else 0.0
        init_num = float(init_bal) if init_bal is not None else 0.0
        add_num = float(added) if added is not None else 0.0
        disb_num = float(disb) if disb is not None else 0.0

        items.append({
            "id": code_str,
            "code": code_str,
            "name": name_str,
            "category": dept_str,
            "unit": unit_str,
            "stock": bal_num,
            "openingBalance": init_num,
            "totalIn": add_num,
            "totalOut": disb_num,
            "alertLimit": 10,
            "classification": "ACTIVE" if bal_num > 0 else "CRITICAL",
            "status": "ACTIVE" if bal_num > 0 else "CRITICAL"
        })
        sum_bal += bal_num

print(f"Total items extracted: {len(items)}")
print(f"Total stock sum: {sum_bal:,.2f}")
assert len(items) == 902, f"Expected 902 items, got {len(items)}"
assert abs(sum_bal - 14733.25) < 0.01, f"Expected 14,733.25, got {sum_bal}"

with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
    json.dump(items, f, ensure_ascii=False, indent=2)

print(f"Successfully generated clean snapshot at: {OUTPUT_JSON}")
