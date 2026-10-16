# -*- coding: utf-8 -*-
import openpyxl
import win32com.client
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'

print(f"Loading {excel_path}...")
wb = openpyxl.load_workbook(excel_path)
ws = wb['الصرف']

fixed_col6 = 0
fixed_col7 = 0
fixed_col11 = 0

for r in range(4, ws.max_row + 1):
    code_val = ws.cell(r, 5).value
    # Skip canceled vouchers (code 9999)
    if code_val is not None and str(code_val).strip() == '9999':
        continue
    
    # Check Col 6 (اسم الصنف)
    c6_val = str(ws.cell(r, 6).value or '')
    if not c6_val.startswith('='):
        ws.cell(r, 6).value = f'=IFERROR(VLOOKUP(E{r},\'أرصدة المخزن\'!$A:$E,3,0),"")'
        fixed_col6 += 1
        
    # Check Col 7 (الوحدة)
    c7_val = str(ws.cell(r, 7).value or '')
    if not c7_val.startswith('='):
        ws.cell(r, 7).value = f'=IFERROR(VLOOKUP(E{r},\'أرصدة المخزن\'!$A:$H,4,FALSE),"")'
        fixed_col7 += 1
        
    # Check Col 11 (الرصيد)
    c11_val = str(ws.cell(r, 11).value or '')
    if not c11_val.startswith('='):
        ws.cell(r, 11).value = f'=IFERROR(VLOOKUP(E{r},\'أرصدة المخزن\'!$A:$H,8,0),"")'
        fixed_col11 += 1

print(f"Restored formulas in الصرف:")
print(f"  Col 6 (اسم الصنف): {fixed_col6} rows restored to dynamic VLOOKUP")
print(f"  Col 7 (الوحدة): {fixed_col7} rows restored to dynamic VLOOKUP")
print(f"  Col 11 (الرصيد): {fixed_col11} rows restored to dynamic VLOOKUP")

wb.save(excel_path)
wb.close()
print("Saved Excel file successfully. Now running Excel COM recalculation...")

# Run COM recalculation and save cached formula values
excel_app = win32com.client.DispatchEx('Excel.Application')
excel_app.Visible = False
excel_app.DisplayAlerts = False
try:
    wb_com = excel_app.Workbooks.Open(excel_path)
    wb_com.Save()
    wb_com.Close(SaveChanges=True)
    print("Excel COM recalculation and save completed successfully!")
finally:
    excel_app.Quit()
