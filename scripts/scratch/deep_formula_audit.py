# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=False)

# 1. Audit أرصدة المخزن formulas
ws_stock = wb['أرصدة المخزن']
print(f"=== Sheet: أرصدة المخزن (Max row: {ws_stock.max_row}) ===")
broken_stock = []
for r in range(4, ws_stock.max_row + 1):
    code = ws_stock.cell(r, 1).value
    name = ws_stock.cell(r, 3).value
    if code is None and name is None:
        continue
    f_form = str(ws_stock.cell(r, 6).value or '')
    g_form = str(ws_stock.cell(r, 7).value or '')
    h_form = str(ws_stock.cell(r, 8).value or '')
    
    if not f_form.startswith('='):
        broken_stock.append((r, code, name, 'Col F (In)', f_form))
    if not g_form.startswith('='):
        broken_stock.append((r, code, name, 'Col G (Out)', g_form))
    if not h_form.startswith('='):
        broken_stock.append((r, code, name, 'Col H (Bal)', h_form))

print(f"Broken formulas in أرصدة المخزن: {len(broken_stock)}")
for b in broken_stock[:10]:
    print(" ", b)

# 2. Audit الاضافة formulas
ws_add = wb['الاضافة']
print(f"\n=== Sheet: الاضافة (Max row: {ws_add.max_row}) ===")
broken_add = []
for r in range(4, ws_add.max_row + 1):
    code = ws_add.cell(r, 4).value
    if code is None or str(code).strip() == '':
        continue
    f_form = str(ws_add.cell(r, 5).value or '') # Name
    g_form = str(ws_add.cell(r, 6).value or '') # Unit
    if not f_form.startswith('='):
        broken_add.append((r, code, 'Col E (Name)', f_form))
    if not g_form.startswith('='):
        broken_add.append((r, code, 'Col F (Unit)', g_form))

print(f"Broken formulas in الاضافة: {len(broken_add)}")
for b in broken_add:
    print(" ", b)

# 3. Audit الصرف formulas
ws_disp = wb['الصرف']
print(f"\n=== Sheet: الصرف (Max row: {ws_disp.max_row}) ===")
broken_disp = []
for r in range(4, ws_disp.max_row + 1):
    code = ws_disp.cell(r, 5).value
    if code is None or str(code).strip() == '' or str(code) == '9999':
        continue
    voucher = ws_disp.cell(r, 4).value
    f_form = str(ws_disp.cell(r, 6).value or '') # Name
    g_form = str(ws_disp.cell(r, 7).value or '') # Unit
    k_form = str(ws_disp.cell(r, 11).value or '') # Stock
    
    if not f_form.startswith('='):
        broken_disp.append((r, voucher, code, 'Col F (Name)', f_form))
    if not g_form.startswith('='):
        broken_disp.append((r, voucher, code, 'Col G (Unit)', g_form))
    if not k_form.startswith('='):
        broken_disp.append((r, voucher, code, 'Col K (Stock)', k_form))

print(f"Broken formulas in الصرف: {len(broken_disp)}")
for b in broken_disp:
    print(f"  Row {b[0]} | Voucher: {b[1]} | Code: {b[2]} | {b[3]}: '{b[4]}'")
