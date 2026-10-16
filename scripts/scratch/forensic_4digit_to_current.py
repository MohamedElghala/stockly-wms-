import openpyxl

ARCHIVE = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"

FILES = {
    "V2.1_2_4DIGIT": ARCHIVE + r"\The_Calm_WMS_V2.1 (2).xlsx",
    "V2.1_3_6DIGIT": ARCHIVE + r"\The_Calm_WMS_V2.1 (3).xlsx",
    "XLSB_SEP13": ARCHIVE + r"\The_Calm_WMS_V2 (version 1).xlsb.xlsx",
    "BEFORE_709": ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx",
    "BEFORE_RESTRUCT": ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx",
    "AFTER_RESTRUCT": ARCHIVE + r"\The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx",
    "CURRENT": r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx",
}

def safe_float(val):
    if val is None or val == "":
        return 0.0
    try:
        return float(val)
    except Exception:
        return 0.0

def load_summary(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    # Compute direct sums from الاضافة and الصرف by resolved item name!
    ws_bal = wb["أرصدة المخزن"]
    cat_by_code = {}
    cat_by_name = {}
    for row in ws_bal.iter_rows(min_row=7, max_col=9, values_only=True):
        code, name, dept, unit, op, ti, to, bal, st = row[:9]
        if code is not None and name is not None and str(code).strip() != "":
            c_str = str(int(code)) if isinstance(code, (int, float)) else str(code).strip()
            n_str = str(name).strip()
            cat_by_code[c_str] = n_str
            if n_str not in cat_by_name:
                cat_by_name[n_str] = {"codes": [], "open": 0.0}
            cat_by_name[n_str]["codes"].append(c_str)
            cat_by_name[n_str]["open"] += safe_float(op)

    in_by_name = {}
    in_total_rows = 0
    in_total_qty = 0.0
    ws_in = wb["الاضافة"]
    for idx, row in enumerate(ws_in.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        seq, date, vno, _, code, name_val, unit, qty, mtype = row[:9]
        if code is None and name_val is None and qty is None:
            continue
        c_str = str(int(code)) if isinstance(code, (int, float)) else (str(code).strip() if code is not None else "")
        res_name = cat_by_code.get(c_str, str(name_val).strip() if name_val else f"UNMAPPED:{c_str}")
        q = safe_float(qty)
        in_by_name[res_name] = in_by_name.get(res_name, 0.0) + q
        in_total_rows += 1
        in_total_qty += q

    out_by_name = {}
    out_total_rows = 0
    out_total_qty = 0.0
    ws_out = wb["الصرف"]
    for idx, row in enumerate(ws_out.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        seq, date, vno, code, name_val, unit, qty, rec, dept = row[:9]
        if code is None and name_val is None and qty is None:
            continue
        c_str = str(int(code)) if isinstance(code, (int, float)) else (str(code).strip() if code is not None else "")
        res_name = cat_by_code.get(c_str, str(name_val).strip() if name_val else f"UNMAPPED:{c_str}")
        q = safe_float(qty)
        out_by_name[res_name] = out_by_name.get(res_name, 0.0) + q
        out_total_rows += 1
        out_total_qty += q

    wb.close()
    return {
        "cat_by_name": cat_by_name,
        "in_by_name": in_by_name,
        "out_by_name": out_by_name,
        "in_total_rows": in_total_rows,
        "in_total_qty": in_total_qty,
        "out_total_rows": out_total_rows,
        "out_total_qty": out_total_qty,
    }

data = {}
for k, p in FILES.items():
    print(f"Loading {k}...")
    data[k] = load_summary(p)
    d = data[k]
    print(f"  {k}: Active Names={len(d['cat_by_name'])}, IN rows={d['in_total_rows']} (qty={d['in_total_qty']}), OUT rows={d['out_total_rows']} (qty={d['out_total_qty']})")

print("\n=== CHECK 1: V2.1_2_4DIGIT vs V2.1_3_6DIGIT (First Code Migration on 2026-09-09) ===")
d1 = data["V2.1_2_4DIGIT"]
d2 = data["V2.1_3_6DIGIT"]
all_names = sorted(set(d1["cat_by_name"].keys()) | set(d2["cat_by_name"].keys()))
diff_12 = 0
for n in all_names:
    op1 = d1["cat_by_name"].get(n, {}).get("open", 0.0)
    op2 = d2["cat_by_name"].get(n, {}).get("open", 0.0)
    in1 = d1["in_by_name"].get(n, 0.0)
    in2 = d2["in_by_name"].get(n, 0.0)
    out1 = d1["out_by_name"].get(n, 0.0)
    out2 = d2["out_by_name"].get(n, 0.0)
    if abs(op1 - op2) > 1e-5 or abs(in1 - in2) > 1e-5 or abs(out1 - out2) > 1e-5:
        diff_12 += 1
        print(f"  {n}: 4DIGIT(op={op1}, in={in1}, out={out1}) -> 6DIGIT(op={op2}, in={in2}, out={out2})")
print(f"Total diffs between 4DIGIT and 6DIGIT: {diff_12}")

print("\n=== CHECK 2: BEFORE_RESTRUCT vs AFTER_RESTRUCT (Second Code Migration on 2026-09-30) ===")
d3 = data["BEFORE_RESTRUCT"]
d4 = data["AFTER_RESTRUCT"]
all_names_34 = sorted(set(d3["cat_by_name"].keys()) | set(d4["cat_by_name"].keys()))
diff_34 = 0
for n in all_names_34:
    op3 = d3["cat_by_name"].get(n, {}).get("open", 0.0)
    op4 = d4["cat_by_name"].get(n, {}).get("open", 0.0)
    in3 = d3["in_by_name"].get(n, 0.0)
    in4 = d4["in_by_name"].get(n, 0.0)
    out3 = d3["out_by_name"].get(n, 0.0)
    out4 = d4["out_by_name"].get(n, 0.0)
    if abs(op3 - op4) > 1e-5 or abs(in3 - in4) > 1e-5 or abs(out3 - out4) > 1e-5:
        diff_34 += 1
        print(f"  {n}: BEFORE(op={op3}, in={in3}, out={out3}) -> AFTER(op={op4}, in={in4}, out={out4})")
print(f"Total diffs between BEFORE_RESTRUCT and AFTER_RESTRUCT: {diff_34}")
