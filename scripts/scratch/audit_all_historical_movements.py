# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

ws_stock = wb['أرصدة المخزن']
ws_add = wb['الاضافة']
ws_disp = wb['الصرف']

print("==================================================")
print("     فحص شامل وتدقيق رياضي لحركات الإكسيل القديمة والحديثة")
print("==================================================")

# 1. Catalog loading
catalog = {} # code -> {name, unit, open, in_sum, out_sum, bal, row}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(r, 1).value
    dept = ws_stock.cell(r, 2).value
    name = ws_stock.cell(r, 3).value
    unit = ws_stock.cell(r, 4).value
    op = ws_stock.cell(r, 5).value or 0
    in_val = ws_stock.cell(r, 6).value or 0
    out_val = ws_stock.cell(r, 7).value or 0
    bal_val = ws_stock.cell(r, 8).value or 0
    
    if c is not None and str(c).strip() != '' and str(c).isdigit():
        code_int = int(c)
        catalog[code_int] = {
            'row': r,
            'dept': dept,
            'name': name,
            'unit': unit,
            'open': float(op),
            'in_sheet': float(in_val),
            'out_sheet': float(out_val),
            'bal_sheet': float(bal_val),
            'calc_in': 0.0,
            'calc_out': 0.0
        }

print(f"1. دليل الأصناف في أرصدة المخزن: {len(catalog)} صنفاً نشطاً.")

# 2. Audit الاضافة
add_errors = []
total_add_qty = 0.0
for r in range(7, 1692):
    code = ws_add.cell(r, 4).value
    name = ws_add.cell(r, 5).value
    unit = ws_add.cell(r, 6).value
    qty = ws_add.cell(r, 7).value
    mov = ws_add.cell(r, 8).value
    
    if code is None or str(code).strip() == '':
        continue
    
    try:
        c_int = int(code)
    except:
        add_errors.append((r, f"كود غير رقمي: {code}"))
        continue
        
    if c_int not in catalog:
        add_errors.append((r, f"كود غير موجود بدليل الأصناف: {c_int}"))
        continue
        
    q_val = float(qty or 0)
    total_add_qty += q_val
    catalog[c_int]['calc_in'] += q_val
    
    # Check evaluated name matching
    if name != catalog[c_int]['name']:
        add_errors.append((r, f"اسم الصنف المقروء '{name}' لا يطابق الدليل '{catalog[c_int]['name']}'"))

print(f"2. حركات الإضافة: {1691 - 6} حركة مفحوصة (إجمالي الكميات المضافة: {total_add_qty:,.2f}) - عدد الأخطاء: {len(add_errors)}")

# 3. Audit الصرف
disp_errors = []
total_disp_qty = 0.0
for r in range(7, 2647):
    voucher = ws_disp.cell(r, 4).value
    code = ws_disp.cell(r, 5).value
    name = ws_disp.cell(r, 6).value
    unit = ws_disp.cell(r, 7).value
    qty = ws_disp.cell(r, 8).value
    mov = ws_disp.cell(r, 9).value
    
    if code is None or str(code).strip() == '':
        continue
    if str(code).strip() == '9999': # Canceled voucher
        continue
        
    try:
        c_int = int(code)
    except:
        disp_errors.append((r, voucher, f"كود غير رقمي: {code}"))
        continue
        
    if c_int not in catalog:
        disp_errors.append((r, voucher, f"كود غير موجود بدليل الأصناف: {c_int}"))
        continue
        
    q_val = float(qty or 0)
    total_disp_qty += q_val
    catalog[c_int]['calc_out'] += q_val
    
    # Check evaluated name matching
    if name != catalog[c_int]['name']:
        disp_errors.append((r, voucher, f"اسم الصنف المقروء '{name}' لا يطابق الدليل '{catalog[c_int]['name']}'"))

print(f"3. حركات الصرف: {2646 - 6} حركة مفحوصة (إجمالي الكميات المنصرفة: {total_disp_qty:,.2f}) - عدد الأخطاء: {len(disp_errors)}")

# 4. Math & Parity Check with أرصدة المخزن
stock_diffs = []
for c_int, it in catalog.items():
    diff_in = abs(it['in_sheet'] - it['calc_in'])
    diff_out = abs(it['out_sheet'] - it['calc_out'])
    expected_bal = it['open'] + it['calc_in'] - it['calc_out']
    diff_bal = abs(it['bal_sheet'] - expected_bal)
    
    if diff_in > 0.001 or diff_out > 0.001 or diff_bal > 0.001:
        stock_diffs.append((c_int, it['name'], it['in_sheet'], it['calc_in'], it['out_sheet'], it['calc_out'], it['bal_sheet'], expected_bal))

print(f"4. مطابقة أرصدة المخزن الرياضية مع دفاتر الإضافة والصرف: {len(catalog)} صنفاً.")
print(f"   عدد الفروقات المكتشفة: {len(stock_diffs)}")
for d in stock_diffs[:10]:
    print(f"   - كود {d[0]} ({d[1]}): وارد شيت={d[2]} مقابل وارد حركات={d[3]} | صرف شيت={d[4]} مقابل صرف حركات={d[5]}")

print("==================================================")
if len(add_errors) == 0 and len(disp_errors) == 0 and len(stock_diffs) == 0:
    print("✅ النتيجة: تطابق جنائي ورياضي بنسبة 100.00% لكافة الحركات القديمة والحديثة، صفر أخطاء وصفر فروقات!")
else:
    print("⚠️ هناك ملاحظات تحتاج تدقيق!")
