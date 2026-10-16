import openpyxl

ARCHIVE = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"

FILES = [
    ("0_V2.1_ORIG_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1.xlsx"),
    ("1_V2.1_1_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (1).xlsx"),
    ("2_V2.1_2_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (2).xlsx"),
    ("3_V2.1_3_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (3).xlsx"),
    ("4_XLSB_SEP13", ARCHIVE + r"\The_Calm_WMS_V2 (version 1).xlsb.xlsx"),
    ("5_BEFORE_709_SEP27", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx"),
    ("6_BEFORE_RESTRUCT_SEP30", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx"),
    ("7_AFTER_RESTRUCT_OCT02", ARCHIVE + r"\The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx"),
    ("8_CURRENT_OCT11", r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"),
]

def safe_float(val, label="", r=0):
    if val is None or val == "":
        return 0.0
    try:
        return float(val)
    except Exception as e:
        return f"ERR({type(val).__name__}:{val})"

def load_wb_data(label, path, out_lines):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = wb.sheetnames
    out_lines.append(f"[{label}] Sheets: {sheets}")
    if "أرصدة المخزن" not in sheets:
        wb.close()
        return None

    ws_bal = wb["أرصدة المخزن"]
    cat_by_code = {}
    cat_by_name = {}
    for idx, row in enumerate(ws_bal.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        code, name, dept, unit, op, ti, to, bal, st = row[:9]
        if code is not None and name is not None and str(code).strip() != "":
            c_str = str(int(code)) if isinstance(code, (int, float)) else str(code).strip()
            n_str = str(name).strip()
            op_val = safe_float(op, label, idx)
            if isinstance(op_val, str):
                out_lines.append(f"  [WEIRD OPEN] {label} row {idx}: {op_val}")
                op_val = 0.0
            cat_by_code[c_str] = {"name": n_str, "dept": str(dept).strip() if dept else "", "open": op_val, "row": idx}
            if n_str not in cat_by_name:
                cat_by_name[n_str] = {"codes": [], "open": 0.0}
            cat_by_name[n_str]["codes"].append(c_str)
            cat_by_name[n_str]["open"] += op_val

    add_rows = {}
    in_by_name = {}
    ws_in = wb["الاضافة"]
    for idx, row in enumerate(ws_in.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        seq, date, vno, _, code, name_val, unit, qty, mtype = row[:9]
        if code is None and name_val is None and qty is None:
            continue
        c_str = str(int(code)) if isinstance(code, (int, float)) else (str(code).strip() if code is not None else "")
        res_name = cat_by_code[c_str]["name"] if c_str in cat_by_code else (str(name_val).strip() if name_val else f"UNMAPPED:{c_str}")
        q_val = safe_float(qty, label, idx)
        if isinstance(q_val, str):
            out_lines.append(f"  [WEIRD ADD QTY] {label} row {idx} (V={vno}, code={c_str}, name={res_name}): {q_val}")
            q_num = 0.0
        else:
            q_num = q_val
        add_rows[idx] = {
            "vno": vno,
            "date": str(date)[:10] if date else "",
            "code": c_str,
            "name": res_name,
            "qty": q_val,
            "qty_num": q_num,
        }
        in_by_name[res_name] = in_by_name.get(res_name, 0.0) + q_num

    dis_rows = {}
    out_by_name = {}
    ws_out = wb["الصرف"]
    for idx, row in enumerate(ws_out.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        seq, date, vno, code, name_val, unit, qty, rec, dept = row[:9]
        if code is None and name_val is None and qty is None:
            continue
        c_str = str(int(code)) if isinstance(code, (int, float)) else (str(code).strip() if code is not None else "")
        res_name = cat_by_code[c_str]["name"] if c_str in cat_by_code else (str(name_val).strip() if name_val else f"UNMAPPED:{c_str}")
        q_val = safe_float(qty, label, idx)
        if isinstance(q_val, str):
            out_lines.append(f"  [WEIRD DIS QTY] {label} row {idx} (V={vno}, code={c_str}, name={res_name}): {q_val}")
            q_num = 0.0
        else:
            q_num = q_val
        dis_rows[idx] = {
            "vno": vno,
            "date": str(date)[:10] if date else "",
            "code": c_str,
            "name": res_name,
            "qty": q_val,
            "qty_num": q_num,
            "rec": str(rec).strip() if rec else "",
        }
        out_by_name[res_name] = out_by_name.get(res_name, 0.0) + q_num

    wb.close()
    return {
        "cat_by_code": cat_by_code,
        "cat_by_name": cat_by_name,
        "add_rows": add_rows,
        "dis_rows": dis_rows,
        "in_by_name": in_by_name,
        "out_by_name": out_by_name,
    }

out_lines = []
loaded = {}
valid_keys = []
for label, path in FILES:
    res = load_wb_data(label, path, out_lines)
    if res is not None:
        loaded[label] = res
        valid_keys.append(label)
        out_lines.append(f"  -> Catalog codes={len(res['cat_by_code'])}, unique names={len(res['cat_by_name'])}, Add rows={len(res['add_rows'])}, Dis rows={len(res['dis_rows'])}")

# Compare consecutive valid versions row-by-row for الاضافة and الصرف (using resolved item name!)
for i in range(len(valid_keys) - 1):
    k1, k2 = valid_keys[i], valid_keys[i+1]
    d1, d2 = loaded[k1], loaded[k2]
    out_lines.append("\n" + "="*80)
    out_lines.append(f"COMPARISON: {k1}  --->  {k2}")
    out_lines.append("="*80)

    # Check common Addition rows where resolved item name or qty changed
    common_add = sorted(set(d1["add_rows"].keys()) & set(d2["add_rows"].keys()))
    add_diffs = []
    for r in common_add:
        r1, r2 = d1["add_rows"][r], d2["add_rows"][r]
        if r1["name"] != r2["name"] or r1["qty"] != r2["qty"] or str(r1["vno"]) != str(r2["vno"]):
            add_diffs.append((r, r1, r2))
    removed_add = sorted(set(d1["add_rows"].keys()) - set(d2["add_rows"].keys()))
    added_add = sorted(set(d2["add_rows"].keys()) - set(d1["add_rows"].keys()))

    out_lines.append(f"[الاضافة] Common rows={len(common_add)}, Modified={len(add_diffs)}, Removed={len(removed_add)}, Added={len(added_add)} (rows {min(added_add) if added_add else '-'}..{max(added_add) if added_add else '-'})")
    for r, r1, r2 in add_diffs:
        out_lines.append(f"  [ADD MOD] Row {r:4d}: V={r1['vno']} [{r1['code']}] '{r1['name']}' (qty={r1['qty']})  -->  V={r2['vno']} [{r2['code']}] '{r2['name']}' (qty={r2['qty']})")
    for r in removed_add:
        r1 = d1["add_rows"][r]
        out_lines.append(f"  [ADD DEL] Row {r:4d}: V={r1['vno']} [{r1['code']}] '{r1['name']}' (qty={r1['qty']})")

    # Check common Disbursement rows where resolved item name or qty changed
    common_dis = sorted(set(d1["dis_rows"].keys()) & set(d2["dis_rows"].keys()))
    dis_diffs = []
    for r in common_dis:
        r1, r2 = d1["dis_rows"][r], d2["dis_rows"][r]
        if r1["name"] != r2["name"] or r1["qty"] != r2["qty"] or str(r1["vno"]) != str(r2["vno"]):
            dis_diffs.append((r, r1, r2))
    removed_dis = sorted(set(d1["dis_rows"].keys()) - set(d2["dis_rows"].keys()))
    added_dis = sorted(set(d2["dis_rows"].keys()) - set(d1["dis_rows"].keys()))

    out_lines.append(f"[الصرف]   Common rows={len(common_dis)}, Modified={len(dis_diffs)}, Removed={len(removed_dis)}, Added={len(added_dis)} (rows {min(added_dis) if added_dis else '-'}..{max(added_dis) if added_dis else '-'})")
    for r, r1, r2 in dis_diffs:
        out_lines.append(f"  [DIS MOD] Row {r:4d}: V={r1['vno']} [{r1['code']}] '{r1['name']}' (qty={r1['qty']})  -->  V={r2['vno']} [{r2['code']}] '{r2['name']}' (qty={r2['qty']})")
    for r in removed_dis:
        r1 = d1["dis_rows"][r]
        out_lines.append(f"  [DIS DEL] Row {r:4d}: V={r1['vno']} [{r1['code']}] '{r1['name']}' (qty={r1['qty']})")

with open(r"c:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\wms-monorepo\scripts\scratch\forensic_report.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out_lines))
print("DONE! Lines written:", len(out_lines))
