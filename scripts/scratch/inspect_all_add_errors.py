# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws_stock = wb['أرصدة المخزن']
ws_add = wb['الاضافة']

catalog = {}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(r, 1).value
    n = ws_stock.cell(r, 3).value
    if c and str(c).isdigit():
        catalog[int(c)] = n

all_errs = []
for r in range(7, 1692):
    code = ws_add.cell(r, 4).value
    name = ws_add.cell(r, 5).value
    if code is None or str(code).strip() == '':
        continue
    try:
        c_int = int(code)
    except:
        all_errs.append((r, f"Non-int code {code}"))
        continue
    if c_int not in catalog:
        all_errs.append((r, f"Code not in catalog {c_int}"))
        continue
    if name != catalog[c_int]:
        all_errs.append((r, f"Row {r} | Code {c_int}: add name '{name}' != catalog '{catalog[c_int]}'"))

print(f"All errors count: {len(all_errs)}")
for e in all_errs:
    print(e)
