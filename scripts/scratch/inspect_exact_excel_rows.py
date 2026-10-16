# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_sarf = wb['الصرف']

# Let's inspect all rows from row 2600 to 2646 (Vouchers 731 to 745)
# Print all columns for each row
headers = [ws_sarf.cell(row=3, column=c).value for c in range(1, 12)]
print("Headers:", headers)

for r in range(2600, 2647):
    row_vals = [ws_sarf.cell(row=r, column=c).value for c in range(1, 12)]
    print(f"Row {r}: {row_vals}")
