# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_add = wb['الاضافة']

non_formula_add = []
for r in range(7, 1692):
    code = ws_add.cell(r, 4).value
    if code and str(code).strip() == '9999':
        continue # Canceled voucher
    
    e_val = str(ws_add.cell(r, 5).value or '')
    f_val = str(ws_add.cell(r, 6).value or '')
    
    if not e_val.startswith('='):
        non_formula_add.append((r, 'Col E (Name)', ws_add.cell(r, 4).value, e_val))
    if not f_val.startswith('='):
        non_formula_add.append((r, 'Col F (Unit)', ws_add.cell(r, 4).value, f_val))

print(f"Non-formula cells in الاضافة (excluding canceled 9999): {len(non_formula_add)}")
for item in non_formula_add:
    print(item)
