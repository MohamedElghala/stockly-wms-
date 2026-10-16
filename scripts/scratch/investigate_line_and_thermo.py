# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws_stock = wb['أرصدة المخزن']
ws_add = wb['الاضافة']
ws_disp = wb['الصرف']

print("==================================================")
print("1. فحص واردات شركة لاين الأخيرة (أذون 272 حتى 276)")
print("==================================================")

line_additions = []
# Vouchers 272 to 276 in الاضافة are rows 1667 to 1691
for r in range(1667, 1692):
    v = ws_add.cell(r, 3).value
    code = ws_add.cell(r, 4).value
    name = ws_add.cell(r, 5).value
    unit = ws_add.cell(r, 6).value
    qty = ws_add.cell(r, 7).value
    supplier = ws_add.cell(r, 9).value
    line_additions.append({
        'row': r,
        'voucher': v,
        'code': code,
        'name': name,
        'unit': unit,
        'in_qty': float(qty or 0),
        'supplier': supplier
    })

# Check current balances of these items in أرصدة المخزن
catalog = {}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(r, 1).value
    if c and str(c).isdigit():
        catalog[int(c)] = {
            'row': r,
            'name': ws_stock.cell(r, 3).value,
            'unit': ws_stock.cell(r, 4).value,
            'open': float(ws_stock.cell(r, 5).value or 0),
            'in': float(ws_stock.cell(r, 6).value or 0),
            'out': float(ws_stock.cell(r, 7).value or 0),
            'bal': float(ws_stock.cell(r, 8).value or 0)
        }

# Group items from line additions
print(f"إجمالي بنود واردات لاين: {len(line_additions)} حركة.")
seen_codes = set()
for item in line_additions:
    code = int(item['code'])
    cat = catalog.get(code, {})
    seen_codes.add(code)

print("\nجدول تفصيلي لواردات لاين الأخيرة وموقف صرفها ورصيدها المتبقي بالمخزن:")
print(f"{'كود':<8} | {'اسم الصنف':<30} | {'الوحدة':<8} | {'وارد لاين':<10} | {'إجمالي الوارد':<12} | {'إجمالي المنصرف':<14} | {'الرصيد المتبقي':<14} | {'حالة الصرف'}")
print("-" * 120)

for item in line_additions:
    code = int(item['code'])
    cat = catalog.get(code, {})
    in_qty = item['in_qty']
    tot_in = cat.get('in', 0)
    tot_out = cat.get('out', 0)
    bal = cat.get('bal', 0)
    
    # Status
    if bal == 0:
        status = "تم صرفه بالكامل (رصيد 0)"
    elif bal == in_qty:
        status = "باقي بالكامل بدون صرف ⚠️"
    elif bal < in_qty:
        status = f"تم صرف جزء ومتبقي {bal:.1f}"
    else:
        status = f"متبقي رصيد بالمخزن {bal:.1f}"
        
    print(f"{code:<8} | {item['name']:<30} | {item['unit']:<8} | {in_qty:<10.1f} | {tot_in:<12.1f} | {tot_out:<14.1f} | {bal:<14.1f} | {status}")

print("\n" + "="*50)
print("2. فحص سلك / كابل 3*3 ترمو والتكرار في الدليل")
print("="*50)

thermo_items = []
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(r, 1).value
    n = str(ws_stock.cell(r, 3).value or '')
    u = str(ws_stock.cell(r, 4).value or '')
    
    if any(k in n for k in ['ترمو', '3*3', '3 * 3', '3×3']):
        thermo_items.append({
            'row': r,
            'code': c,
            'dept': ws_stock.cell(r, 2).value,
            'name': n,
            'unit': u,
            'open': ws_stock.cell(r, 5).value,
            'in': ws_stock.cell(r, 6).value,
            'out': ws_stock.cell(r, 7).value,
            'bal': ws_stock.cell(r, 8).value
        })

print(f"الأصناف التي تحتوي على 'ترمو' أو '3*3' في أرصدة المخزن: {len(thermo_items)}")
for it in thermo_items:
    print(f"صف {it['row']} | كود: {it['code']} | قسم: {it['dept']} | اسم: {it['name']} | وحدة: {it['unit']} | رصيد أول: {it['open']} | وارد: {it['in']} | منصرف: {it['out']} | رصيد: {it['bal']}")

# Also search for any wire / cable items in electrical department
print("\nفحص كافة أصناف السلك في قسم الأدوات الكهربائية للتحقق من أي تكرار:")
wire_items = []
for r in range(4, 504):
    c = ws_stock.cell(r, 1).value
    n = str(ws_stock.cell(r, 3).value or '')
    u = str(ws_stock.cell(r, 4).value or '')
    if any(k in n for k in ['سلك', 'كابل', 'ترمو', '3*3']):
        wire_items.append((r, c, n, u, ws_stock.cell(r, 8).value))

for w in wire_items:
    print(f"  صف {w[0]:3d} | كود {w[1]} | {w[2]} | وحدة: {w[3]} | رصيد: {w[4]}")

# Also check movements of code 100109 in الاضافة and الصرف
print("\nحركات صنف كود 100109 (لفة سلك 100 م ترمو 3*3):")
print("--- في الاضافة ---")
for r in range(7, 1692):
    if ws_add.cell(r, 4).value == 100109 or str(ws_add.cell(r, 4).value) == '100109':
        print(f"  صف {r} | إذن {ws_add.cell(r, 3).value} | تاريخ: {ws_add.cell(r, 2).value} | كمية: {ws_add.cell(r, 7).value} | ملاحظات: {ws_add.cell(r, 9).value}")

print("--- في الصرف ---")
for r in range(7, 2647):
    if ws_disp.cell(r, 5).value == 100109 or str(ws_disp.cell(r, 5).value) == '100109':
        print(f"  صف {r} | إذن {ws_disp.cell(r, 4).value} | تاريخ: {ws_disp.cell(r, 2).value} | كمية: {ws_disp.cell(r, 8).value} | ملاحظات: {ws_disp.cell(r, 10).value}")
