import openpyxl

ARCHIVE = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"

FILES = [
    ("2_V2.1_2_SEP09_4DIGIT", ARCHIVE + r"\The_Calm_WMS_V2.1 (2).xlsx", "4digit"),
    ("3_V2.1_3_SEP09_6DIGIT", ARCHIVE + r"\The_Calm_WMS_V2.1 (3).xlsx", "6digit"),
    ("4_XLSB_SEP13", ARCHIVE + r"\The_Calm_WMS_V2 (version 1).xlsb.xlsx", "6digit"),
    ("5_BEFORE_709_SEP27", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx", "6digit"),
    ("6_BEFORE_RESTRUCT_SEP30", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx", "6digit"),
    ("7_AFTER_RESTRUCT_OCT02", ARCHIVE + r"\The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx", "6digit"),
    ("8_CURRENT_OCT11", r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx", "6digit"),
]

def norm_code(c):
    if c is None or str(c).strip() == "":
        return ""
    if isinstance(c, (int, float)):
        return str(int(round(c)))
    s = str(c).strip()
    try:
        return str(int(round(float(s))))
    except Exception:
        return s

def safe_float(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except Exception:
        return f"ERR({type(v).__name__}:{v})"

def load_version(label, path, mode, log):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    cat_by_code = {}
    cat_by_name = {}
    add_rows = {}
    dis_rows = {}
    in_by_name = {}
    out_by_name = {}

    if mode == "4digit":
        ws_bal = wb["الرصيد الحالي"]
        for r_idx, row in enumerate(ws_bal.iter_rows(min_row=6, max_col=10, values_only=True), start=6):
            code = norm_code(row[1])
            dept = str(row[2]).strip() if row[2] else ""
            name = str(row[3]).strip() if row[3] else ""
            unit = str(row[4]).strip() if row[4] else ""
            op = safe_float(row[5])
            if isinstance(op, str):
                log.append(f"  [WEIRD OPEN] {label} row {r_idx}: {op}")
                op = 0.0
            if code and name:
                cat_by_code[code] = {"name": name, "dept": dept, "unit": unit, "open": op, "row": r_idx}
                if name not in cat_by_name:
                    cat_by_name[name] = {"codes": [], "open": 0.0}
                cat_by_name[name]["codes"].append(code)
                cat_by_name[name]["open"] += op

        ws_in = wb["الإضافة"]
        for r_idx, row in enumerate(ws_in.iter_rows(min_row=5, max_col=12, values_only=True), start=5):
            seq, date, vno, code_raw, name_raw, unit_raw, qty_raw = row[2:9]
            if code_raw is None and name_raw is None and qty_raw is None:
                continue
            c = norm_code(code_raw)
            name = cat_by_code[c]["name"] if c in cat_by_code else (str(name_raw).strip() if name_raw else f"UNMAPPED:{c}")
            q = safe_float(qty_raw)
            q_num = 0.0 if isinstance(q, str) else q
            if isinstance(q, str):
                log.append(f"  [WEIRD ADD QTY] {label} row {r_idx} (seq={seq}, V={vno}, code={c}, name={name}): {q}")
            seq_key = int(round(float(seq))) if isinstance(seq, (int, float)) else (r_idx - 4)
            add_rows[seq_key] = {"row": r_idx, "vno": vno, "code": c, "name": name, "qty": q, "qty_num": q_num}
            in_by_name[name] = in_by_name.get(name, 0.0) + q_num

        ws_out = wb["الصرف"]
        for r_idx, row in enumerate(ws_out.iter_rows(min_row=5, max_col=12, values_only=True), start=5):
            seq, date, vno, rec_raw, code_raw, name_raw, unit_raw, qty_raw = row[0:8]
            if code_raw is None and name_raw is None and qty_raw is None:
                continue
            c = norm_code(code_raw)
            name = cat_by_code[c]["name"] if c in cat_by_code else (str(name_raw).strip() if name_raw else f"UNMAPPED:{c}")
            q = safe_float(qty_raw)
            q_num = 0.0 if isinstance(q, str) else q
            if isinstance(q, str):
                log.append(f"  [WEIRD DIS QTY] {label} row {r_idx} (seq={seq}, V={vno}, code={c}, name={name}): {q}")
            seq_key = int(round(float(seq))) if isinstance(seq, (int, float)) else (r_idx - 4)
            dis_rows[seq_key] = {"row": r_idx, "vno": vno, "code": c, "name": name, "qty": q, "qty_num": q_num, "rec": str(rec_raw).strip() if rec_raw else ""}
            out_by_name[name] = out_by_name.get(name, 0.0) + q_num

    else:
        ws_bal = wb["أرصدة المخزن"]
        for r_idx, row in enumerate(ws_bal.iter_rows(min_row=4, max_col=9, values_only=True), start=4):
            code = norm_code(row[0])
            dept = str(row[1]).strip() if row[1] else ""
            name = str(row[2]).strip() if row[2] else ""
            unit = str(row[3]).strip() if row[3] else ""
            op = safe_float(row[4])
            if isinstance(op, str):
                log.append(f"  [WEIRD OPEN] {label} row {r_idx}: {op}")
                op = 0.0
            if code and name:
                cat_by_code[code] = {"name": name, "dept": dept, "unit": unit, "open": op, "row": r_idx}
                if name not in cat_by_name:
                    cat_by_name[name] = {"codes": [], "open": 0.0}
                cat_by_name[name]["codes"].append(code)
                cat_by_name[name]["open"] += op

        ws_in = wb["الاضافة"]
        for r_idx, row in enumerate(ws_in.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
            seq, date, vno, code_raw, name_raw, unit_raw, qty_raw, mtype, notes = row[:9]
            if code_raw is None and name_raw is None and qty_raw is None:
                continue
            c = norm_code(code_raw)
            name = cat_by_code[c]["name"] if c in cat_by_code else (str(name_raw).strip() if name_raw else f"UNMAPPED:{c}")
            q = safe_float(qty_raw)
            q_num = 0.0 if isinstance(q, str) else q
            if isinstance(q, str):
                log.append(f"  [WEIRD ADD QTY] {label} row {r_idx} (seq={seq}, V={vno}, code={c}, name={name}): {q}")
            seq_key = r_idx - 6
            add_rows[seq_key] = {"row": r_idx, "vno": vno, "code": c, "name": name, "qty": q, "qty_num": q_num}
            in_by_name[name] = in_by_name.get(name, 0.0) + q_num

        ws_out = wb["الصرف"]
        for r_idx, row in enumerate(ws_out.iter_rows(min_row=7, max_col=10, values_only=True), start=7):
            seq, date, rec_raw, vno, code_raw, name_raw, unit_raw, qty_raw, mtype, notes = row[:10]
            if code_raw is None and name_raw is None and qty_raw is None:
                continue
            c = norm_code(code_raw)
            name = cat_by_code[c]["name"] if c in cat_by_code else (str(name_raw).strip() if name_raw else f"UNMAPPED:{c}")
            q = safe_float(qty_raw)
            q_num = 0.0 if isinstance(q, str) else q
            if isinstance(q, str):
                log.append(f"  [WEIRD DIS QTY] {label} row {r_idx} (seq={seq}, V={vno}, code={c}, name={name}): {q}")
            seq_key = r_idx - 6
            dis_rows[seq_key] = {"row": r_idx, "vno": vno, "code": c, "name": name, "qty": q, "qty_num": q_num, "rec": str(rec_raw).strip() if rec_raw else ""}
            out_by_name[name] = out_by_name.get(name, 0.0) + q_num

    wb.close()
    return {
        "cat_by_code": cat_by_code,
        "cat_by_name": cat_by_name,
        "add_rows": add_rows,
        "dis_rows": dis_rows,
        "in_by_name": in_by_name,
        "out_by_name": out_by_name,
    }

log = []
loaded = {}
keys = []
for label, path, mode in FILES:
    d = load_version(label, path, mode, log)
    loaded[label] = d
    keys.append(label)
    tot_open = sum(v["open"] for v in d["cat_by_name"].values())
    tot_in = sum(d["in_by_name"].values())
    tot_out = sum(d["out_by_name"].values())
    log.append(f"[{label}] Active Codes={len(d['cat_by_code'])}, Unique Names={len(d['cat_by_name'])}, Open={tot_open}, Add Rows={len(d['add_rows'])} (Sum={tot_in}), Dis Rows={len(d['dis_rows'])} (Sum={tot_out})")

# Consecutive row-by-row comparisons (by seq_key!)
for i in range(len(keys) - 1):
    k1, k2 = keys[i], keys[i+1]
    d1, d2 = loaded[k1], loaded[k2]
    log.append("\n" + "="*90)
    log.append(f"STEP {i+1}: {k1}  --->  {k2}")
    log.append("="*90)

    # 1. Addition rows
    c_add = sorted(set(d1["add_rows"].keys()) & set(d2["add_rows"].keys()))
    add_mods = []
    for s in c_add:
        r1, r2 = d1["add_rows"][s], d2["add_rows"][s]
        if r1["name"] != r2["name"] or r1["qty"] != r2["qty"]:
            add_mods.append((s, r1, r2))
    rem_add = sorted(set(d1["add_rows"].keys()) - set(d2["add_rows"].keys()))
    new_add = sorted(set(d2["add_rows"].keys()) - set(d1["add_rows"].keys()))

    log.append(f"[الاضافة] Common={len(c_add)}, Modified(Name/Qty)={len(add_mods)}, Removed={len(rem_add)}, New={len(new_add)}")
    for s, r1, r2 in add_mods:
        log.append(f"  [ADD MOD] Seq {s:4d} (Row {r2['row']:4d}, V={r1['vno']}->{r2['vno']}): [{r1['code']}] '{r1['name']}' (qty={r1['qty']})  -->  [{r2['code']}] '{r2['name']}' (qty={r2['qty']})")
    for s in rem_add:
        r1 = d1["add_rows"][s]
        log.append(f"  [ADD DEL] Seq {s:4d} (Row {r1['row']:4d}, V={r1['vno']}): [{r1['code']}] '{r1['name']}' (qty={r1['qty']})")

    # 2. Disbursement rows
    c_dis = sorted(set(d1["dis_rows"].keys()) & set(d2["dis_rows"].keys()))
    dis_mods = []
    for s in c_dis:
        r1, r2 = d1["dis_rows"][s], d2["dis_rows"][s]
        if r1["name"] != r2["name"] or r1["qty"] != r2["qty"]:
            dis_mods.append((s, r1, r2))
    rem_dis = sorted(set(d1["dis_rows"].keys()) - set(d2["dis_rows"].keys()))
    new_dis = sorted(set(d2["dis_rows"].keys()) - set(d1["dis_rows"].keys()))

    log.append(f"[الصرف]   Common={len(c_dis)}, Modified(Name/Qty)={len(dis_mods)}, Removed={len(rem_dis)}, New={len(new_dis)}")
    for s, r1, r2 in dis_mods:
        log.append(f"  [DIS MOD] Seq {s:4d} (Row {r2['row']:4d}, V={r1['vno']}->{r2['vno']}): [{r1['code']}] '{r1['name']}' (qty={r1['qty']})  -->  [{r2['code']}] '{r2['name']}' (qty={r2['qty']})")
    for s in rem_dis:
        r1 = d1["dis_rows"][s]
        log.append(f"  [DIS DEL] Seq {s:4d} (Row {r1['row']:4d}, V={r1['vno']}): [{r1['code']}] '{r1['name']}' (qty={r1['qty']})")

out_path = r"c:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\wms-monorepo\scripts\scratch\forensic_exact_report.txt"
with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(log))
print("DONE! Total lines:", len(log))
