# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws_stock = wb['أرصدة المخزن']
ws_add = wb['الاضافة']

# Catalog
catalog = {}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(r, 1).value
    n = ws_stock.cell(r, 3).value
    if c and str(c).isdigit():
        catalog[int(c)] = (n, r)

print("--- The 9 errors in الاضافة ---")
for r in range(7, 1692):
    code = ws_add.cell(r, 4).value
    name = ws_add.cell(r, 5).value
    if code and str(code).isdigit():
        c_int = int(code)
        if c_int in catalog:
            cat_name = catalog[c_int][0]
            if name != cat_name:
                print(f"Row {r} | Voucher {ws_add.cell(r,3).value} | Code {c_int} | Add Sheet Name: '{name}' vs Catalog: '{cat_name}'")

print("\n--- Code 101334 in أرصدة المخزن ---")
row_101334 = catalog[101334][1]
wb_form = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_stock_f = wb_form['أرصدة المخزن']
print(f"Row {row_101334} values (data_only):", [ws_stock.cell(row_101334, c).value for c in range(1, 9)])
print(f"Row {row_101334} formulas:", [ws_stock_f.cell(row_101334, c).value for c in range(1, 9)])

# Inspect movement for 101334 in الاضافة
print("\n--- Movements for 101334 in الاضافة ---")
for r in range(7, 1692):
    code = ws_add.cell(r, 4).value
    if code and str(code).strip() == '101334':
        print(f"Add Row {r} | V:{ws_add.cell(r,3).value} | Code:{code} | Name:{ws_add.cell(r,5).value} | Qty:{ws_add.cell(r,7).value} | Mov:{ws_add.cell(r,8).value}")
