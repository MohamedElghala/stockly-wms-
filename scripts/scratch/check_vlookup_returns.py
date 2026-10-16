# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_sarf = wb['الصرف']
ws_stock = wb['أرصدة المخزن']

# Build catalog
catalog = {}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(row=r, column=1).value
    n = ws_stock.cell(row=r, column=3).value
    u = ws_stock.cell(row=r, column=4).value
    if c is not None:
        try:
            catalog[int(c)] = (n, u)
        except:
            catalog[c] = (n, u)

print("Checking what VLOOKUP would return for rows 2499 to 2646 in الصرف:")
mismatches = []
for r in range(2499, 2647):
    voucher = ws_sarf.cell(row=r, column=4).value
    code = ws_sarf.cell(row=r, column=5).value
    cell_name = ws_sarf.cell(row=r, column=6).value
    cell_unit = ws_sarf.cell(row=r, column=7).value
    
    # lookup
    cat_item = catalog.get(code)
    if cat_item is None:
        print(f"ERROR: Code {code} at row {r} (Voucher {voucher}) NOT FOUND in أرصدة المخزن!")
    else:
        vlookup_name, vlookup_unit = cat_item
        # If cell_name is static text, does it match vlookup_name?
        if not str(cell_name).startswith('='):
            if cell_name != vlookup_name:
                mismatches.append((r, voucher, code, cell_name, vlookup_name))

print(f"\nMismatches between static text and catalog name: {len(mismatches)}")
for m in mismatches:
    print(f"Row {m[0]} | Voucher {m[1]} | Code {m[2]} | Static: '{m[3]}' vs Catalog: '{m[4]}'")
