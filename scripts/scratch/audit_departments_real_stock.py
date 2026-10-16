import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws = wb['أرصدة المخزن']

departments = {}

for r in range(4, ws.max_row+1):
    c_code = ws.cell(row=r, column=1).value
    c_dept = ws.cell(row=r, column=2).value
    c_name = ws.cell(row=r, column=3).value
    c_unit = str(ws.cell(row=r, column=4).value or 'عدد').strip()
    c_init = float(ws.cell(row=r, column=5).value or 0)
    c_in = float(ws.cell(row=r, column=6).value or 0)
    c_out = float(ws.cell(row=r, column=7).value or 0)
    c_bal = float(ws.cell(row=r, column=8).value or 0)
    
    if c_name and str(c_name).strip():
        dept_name = str(c_dept or 'غير محدد').strip()
        if dept_name not in departments:
            departments[dept_name] = {
                'items_count': 0,
                'items_with_stock': 0,
                'total_init': 0.0,
                'total_in': 0.0,
                'total_out': 0.0,
                'total_bal': 0.0,
                'items': []
            }
        d = departments[dept_name]
        d['items_count'] += 1
        if c_bal > 0:
            d['items_with_stock'] += 1
        d['total_init'] += c_init
        d['total_in'] += c_in
        d['total_out'] += c_out
        d['total_bal'] += c_bal
        d['items'].append({
            'code': str(c_code),
            'name': str(c_name).strip(),
            'unit': c_unit,
            'init': c_init,
            'in': c_in,
            'out': c_out,
            'bal': c_bal
        })

print("================================================================================")
print("             تقرير الفحص الجنائي لأرصدة الأقسام الـ 17 المعتمدة")
print("================================================================================")
grand_bal = sum(d['total_bal'] for d in departments.values())
grand_items = sum(d['items_count'] for d in departments.values())
grand_init = sum(d['total_init'] for d in departments.values())
grand_in = sum(d['total_in'] for d in departments.values())
grand_out = sum(d['total_out'] for d in departments.values())

print(f"إجمالي الأقسام: {len(departments)} قسم")
print(f"إجمالي الأصناف بالكتالوج: {grand_items} صنف")
print(f"إجمالي الرصيد الافتتاحي: {grand_init:,.2f}")
print(f"إجمالي الوارد التاريخي: {grand_in:,.2f}")
print(f"إجمالي المنصرف التاريخي: {grand_out:,.2f}")
print(f"إجمالي الرصيد الحالي بالمخزن: {grand_bal:,.2f} وحدة")
print(f"معادلة التحقق الرياضي: {grand_init:,.2f} + {grand_in:,.2f} - {grand_out:,.2f} = {grand_init + grand_in - grand_out:,.2f}")

print("\n--------------------------------------------------------------------------------")
print(f"{'القسم':<25} | {'أصناف':<5} | {'برصيد':<5} | {'افتتاحي':<10} | {'وارد':<10} | {'منصرف':<10} | {'الرصيد الحالي':<12}")
print("--------------------------------------------------------------------------------")

for d_name, data in sorted(departments.items(), key=lambda x: x[1]['total_bal'], reverse=True):
    print(f"{d_name:<25} | {data['items_count']:<5} | {data['items_with_stock']:<5} | {data['total_init']:<10,.0f} | {data['total_in']:<10,.0f} | {data['total_out']:<10,.0f} | {data['total_bal']:<12,.1f}")

print("\n================================================================================")
print("           تفاصيل كل قسم وأعلى الأصناف رصيداً فيه وكيف تشكلت أرقامه")
print("================================================================================")

for d_name, data in sorted(departments.items(), key=lambda x: x[1]['total_bal'], reverse=True):
    print(f"\n📁 [{d_name}] -> إجمالي الرصيد: {data['total_bal']:,.1f} (من {data['items_count']} صنف، {data['items_with_stock']} منها به رصيد)")
    top_items = sorted(data['items'], key=lambda x: x['bal'], reverse=True)[:3]
    for idx, it in enumerate(top_items, 1):
        if it['bal'] > 0:
            print(f"   {idx}. كود {it['code']}: {it['name']} | الرصيد: {it['bal']:,.1f} {it['unit']} [افتتاحي: {it['init']}, وارد: {it['in']}, منصرف: {it['out']}]")
        else:
            print(f"   (لا توجد أصناف بها رصيد في هذا القسم - الرصيد صفري)")
            break
