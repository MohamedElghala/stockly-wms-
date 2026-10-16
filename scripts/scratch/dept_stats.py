import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx', data_only=True)
ws = wb['أرصدة المخزن']
depts = {}
for r in range(4, ws.max_row + 1):
    dept = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    bal = ws.cell(row=r, column=8).value
    if name and str(name).strip() and dept and str(dept).strip():
        d = str(dept).strip()
        if d not in depts:
            depts[d] = {'total': 0, 'pos': 0, 'zero': 0, 'stock': 0.0}
        depts[d]['total'] += 1
        b = float(bal) if bal is not None else 0.0
        depts[d]['stock'] += b
        if b > 0:
            depts[d]['pos'] += 1
        else:
            depts[d]['zero'] += 1

print(f"{'Idx':3s} | {'Department':30s} | {'Total':5s} | {'Pos':5s} | {'Zero':5s} | {'Total Stock':12s}")
print('-' * 75)
idx = 1
tot_all = 0
tot_pos = 0
tot_zero = 0
tot_stk = 0.0
for d, v in depts.items():
    print(f"{idx:3d} | {d:30s} | {v['total']:5d} | {v['pos']:5d} | {v['zero']:5d} | {v['stock']:12.2f}")
    idx += 1
    tot_all += v['total']
    tot_pos += v['pos']
    tot_zero += v['zero']
    tot_stk += v['stock']
print('-' * 75)
print(f"{'TOT':3s} | {'17 Departments':30s} | {tot_all:5d} | {tot_pos:5d} | {tot_zero:5d} | {tot_stk:12.2f}")
