import openpyxl

ARCHIVE = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"

FILES = [
    ("0_V2.1_ORIG_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1.xlsx"),
    ("1_V2.1_1_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (1).xlsx"),
    ("2_V2.1_2_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (2).xlsx"),
    ("3_V2.1_3_SEP09", ARCHIVE + r"\The_Calm_WMS_V2.1 (3).xlsx"),
    ("4_XLSB_SEP13", ARCHIVE + r"\The_Calm_WMS_V2 (version 1).xlsb.xlsx"),
    ("5_BEFORE_709_SEP27", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_DISBURSEMENT_709.xlsx"),
    ("6_BEFORE_RESTRUCT_SEP30", ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx"),
    ("7_AFTER_RESTRUCT_OCT02", ARCHIVE + r"\The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx"),
    ("8_CURRENT_OCT11", r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"),
]

lines = []
for label, path in FILES:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines.append("="*80)
    lines.append(f"FILE: {label} -> {wb.sheetnames}")
    for sname in wb.sheetnames:
        if sname in ["أرصدة المخزن", "الرصيد الحالي", "الاضافة", "الإضافة", "الصرف"]:
            ws = wb[sname]
            lines.append(f"  Sheet: {sname}")
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, max_col=12, values_only=True), start=1):
                if any(x is not None for x in row):
                    lines.append(f"    Row {r_idx:2d}: {list(row)}")
    wb.close()

with open(r"c:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\wms-monorepo\scripts\scratch\headers_report.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("Headers written:", len(lines))
