import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding="utf-8")

MASTER_EXCEL = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"
AUDIT_EXCEL = r"C:\Users\Mohamed Helmy\Downloads\كشف_جرد_المخزن_حسب_الأقسام.xlsx"
MD_OUT = r"C:\Users\Mohamed Helmy\.gemini\antigravity\brain\4644325f-c2d1-44e5-8356-23b4b57cbdd3\inventory_audit_by_department.md"

wb_src = openpyxl.load_workbook(MASTER_EXCEL, data_only=True, read_only=True)
ws_src = wb_src["أرصدة المخزن"]

depts = {}
dept_order = []
curr_dept = ""

for r_idx, row in enumerate(ws_src.iter_rows(min_row=4, max_row=8503, values_only=True), 4):
    code, dept, name, unit, op, in_q, out_q, bal = (list(row) + [None]*8)[:8]
    if dept and str(dept).strip():
        curr_dept = str(dept).strip()
    if name and str(name).strip():
        if curr_dept not in depts:
            depts[curr_dept] = []
            dept_order.append(curr_dept)
        b_val = float(bal or 0)
        if b_val == int(b_val):
            b_val = int(b_val)
        depts[curr_dept].append({
            "name": str(name).strip(),
            "bal": b_val,
        })
wb_src.close()

wb_out = openpyxl.Workbook()
wb_out.remove(wb_out.active)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(name="Cairo", size=12, bold=True, color="FFFFFF")
bold_font = Font(name="Cairo", size=11, bold=True)
reg_font = Font(name="Cairo", size=11)
zero_fill = PatternFill("solid", fgColor="F9F9F9")
pos_fill = PatternFill("solid", fgColor="E2EFDA")
neg_fill = PatternFill("solid", fgColor="FCE4D6")
thin = Side(border_style="thin", color="808080")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
right_al = Alignment(horizontal="right", vertical="center")

md_lines = [
    "# كشف جرد المخزن لكل قسم لوحده (4 خانات فقط للجرد والتسوية)",
    "",
    "- **ملف الإكسيل الجاهز للطباعة والجرد اليدوي:** [كشف_جرد_المخزن_حسب_الأقسام.xlsx](file:///C:/Users/Mohamed%20Helmy/Downloads/كشف_جرد_المخزن_حسب_الأقسام.xlsx)",
    "",
    "---",
    ""
]

for idx, d_name in enumerate(dept_order, 1):
    safe_title = d_name[:31].replace("/", "-").replace("\\", "-")
    ws = wb_out.create_sheet(title=safe_title)
    ws.views.sheetView[0].rightToLeft = True

    # ONLY the 4 requested columns: اسم الصنف | الرصيد الموجود | صرف للتسوية | إضافة للتسوية
    headers = ["اسم الصنف", "الرصيد الموجود", "صرف للتسوية", "إضافة للتسوية"]
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, 5):
        c = ws.cell(1, col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    # Put items with non-zero balance first, then zero balance items
    d_items = depts[d_name]
    non_zero = [x for x in d_items if x["bal"] != 0]
    zeros = [x for x in d_items if x["bal"] == 0]
    ordered_items = non_zero + zeros

    for i, it in enumerate(ordered_items, 1):
        r_num = i + 1
        ws.append([it["name"], it["bal"], None, None])
        ws.row_dimensions[r_num].height = 24
        row_fill = pos_fill if it["bal"] > 0 else (neg_fill if it["bal"] < 0 else zero_fill)
        for col_idx in range(1, 5):
            c = ws.cell(r_num, col_idx)
            c.font = bold_font if col_idx in (1, 2) else reg_font
            c.alignment = right_al if col_idx == 1 else center
            c.border = border
            if col_idx == 2:
                c.fill = row_fill

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22

    # Page setup for clean A4 printing
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    md_lines.append(f"## {idx}. قسم: {d_name} ({len(non_zero)} صنف له رصيد / إجمالي {len(ordered_items)} صنف)")
    md_lines.append("")
    md_lines.append("| اسم الصنف | الرصيد الموجود | صرف للتسوية | إضافة للتسوية |")
    md_lines.append("| :--- | :-: | :-: | :-: |")
    for it in non_zero:
        md_lines.append(f"| {it['name']} | **`{it['bal']}`** |  |  |")
    md_lines.append("")
    if zeros:
        md_lines.append(f"<details><summary>الأصناف المصفر رصيدها في {d_name} ({len(zeros)} صنف)</summary>\n")
        md_lines.append("| اسم الصنف | الرصيد الموجود | صرف للتسوية | إضافة للتسوية |")
        md_lines.append("| :--- | :-: | :-: | :-: |")
        for it in zeros:
            md_lines.append(f"| {it['name']} | `0` |  |  |")
        md_lines.append("\n</details>")
    md_lines.append("\n---\n")

wb_out.save(AUDIT_EXCEL)
with open(MD_OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines))
print(f"Successfully saved 4-column audit workbook and markdown!")
