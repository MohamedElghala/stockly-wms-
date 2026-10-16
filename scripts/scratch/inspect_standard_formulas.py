# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_sarf = wb['الصرف']

print("=== Standard formulas in الصرف ===")
for r in [10, 50, 100, 500, 1000, 2000, 2499, 2500, 2600, 2614, 2646]:
    print(f"\nRow {r}:")
    for c in range(1, 12):
        val = ws_sarf.cell(row=r, column=c).value
        print(f"  Col {c} ({openpyxl.utils.get_column_letter(c)}): {val}")
