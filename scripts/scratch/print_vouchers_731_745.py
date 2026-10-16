# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=True)
ws = wb['الصرف']

print("--- VOUCHERS 731 to 745 in الصرف ---")
current_v = None
for r in range(2600, 2647):
    v = ws.cell(row=r, column=4).value
    dt = ws.cell(row=r, column=2).value
    dt_str = dt.strftime("%Y-%m-%d") if hasattr(dt, 'strftime') else str(dt)
    dept = ws.cell(row=r, column=3).value
    
    if v != current_v:
        current_v = v
        print(f"\n=== إذن صرف رقم [{v}] (تاريخ: {dt_str}) | الجهة: {dept} ===")
    
    code = ws.cell(row=r, column=5).value
    name = ws.cell(row=r, column=6).value
    unit = ws.cell(row=r, column=7).value
    qty = ws.cell(row=r, column=8).value
    notes = ws.cell(row=r, column=10).value
    print(f"  [صف {r}] كود {code} | {name} | {qty} {unit} | مكان الصرف: {notes}")
