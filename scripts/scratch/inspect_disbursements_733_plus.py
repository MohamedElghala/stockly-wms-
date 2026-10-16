# -*- coding: utf-8 -*-
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx', data_only=False)
ws_sarf = wb['الصرف']
ws_stock = wb['أرصدة المخزن']

# Build catalog from أرصدة المخزن
catalog = {}
for r in range(4, ws_stock.max_row + 1):
    c = ws_stock.cell(row=r, column=1).value
    n = ws_stock.cell(row=r, column=3).value
    u = ws_stock.cell(row=r, column=4).value
    if c is not None:
        try:
            catalog[int(c)] = (n, u, r)
        except:
            catalog[c] = (n, u, r)

print("=== Dissecting rows 2600 to 2646 in الصرف ===")
for r in range(2600, 2647):
    date_val = ws_sarf.cell(row=r, column=2).value
    entity = ws_sarf.cell(row=r, column=3).value
    voucher = ws_sarf.cell(row=r, column=4).value
    code = ws_sarf.cell(row=r, column=5).value
    name = ws_sarf.cell(row=r, column=6).value
    unit = ws_sarf.cell(row=r, column=7).value
    qty = ws_sarf.cell(row=r, column=8).value
    mov = ws_sarf.cell(row=r, column=9).value
    notes = ws_sarf.cell(row=r, column=10).value
    
    cat_item = catalog.get(code, ('NOT FOUND', 'NOT FOUND', 0))
    print(f"Row {r:4d} | V:{voucher} | Code:{code} | Name Cell:{str(name)[:25]} | Catalog Name:{str(cat_item[0])[:25]} | Unit:{unit} | Qty:{qty} | Notes:{notes}")
