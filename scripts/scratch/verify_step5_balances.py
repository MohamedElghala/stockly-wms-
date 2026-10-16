import openpyxl
from collections import defaultdict

ARCHIVE = r"D:\The_Calm_WMS_V2.1\Excel_Database\Historical_Backups_Archive"
P_BEFORE = ARCHIVE + r"\The_Calm_WMS_V2_BACKUP_BEFORE_CATALOG_RESTRUCTURE.xlsx"
P_AFTER  = ARCHIVE + r"\The_Calm_WMS_V2_النسخة_النهائية_المضبوطة.xlsx"

def norm_code(c):
    if c is None or str(c).strip() == "":
        return ""
    if isinstance(c, (int, float)):
        return str(int(round(c)))
    s = str(c).strip()
    try:
        return str(int(round(float(s))))
    except Exception:
        return s

def sfloat(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except Exception:
        return 0.0

def load_state(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws_bal = wb["أرصدة المخزن"]
    cat = {}
    for r_idx, row in enumerate(ws_bal.iter_rows(min_row=4, max_col=9, values_only=True), start=4):
        c = norm_code(row[0])
        dept = str(row[1]).strip() if row[1] else ""
        name = str(row[2]).strip() if row[2] else ""
        unit = str(row[3]).strip() if row[3] else ""
        op = sfloat(row[4])
        if c and name:
            cat[c] = {"name": name, "dept": dept, "unit": unit, "open": op, "in": 0.0, "out": 0.0}

    ws_in = wb["الاضافة"]
    in_rows = {}
    for r_idx, row in enumerate(ws_in.iter_rows(min_row=7, max_col=9, values_only=True), start=7):
        seq, date, vno, code_raw, name_raw, unit_raw, qty_raw = row[:7]
        if code_raw is None and name_raw is None and qty_raw is None:
            continue
        c = norm_code(code_raw)
        q = sfloat(qty_raw)
        in_rows[r_idx] = (c, q)
        if c in cat:
            cat[c]["in"] += q

    ws_out = wb["الصرف"]
    out_rows = {}
    for r_idx, row in enumerate(ws_out.iter_rows(min_row=7, max_col=10, values_only=True), start=7):
        seq, date, rec_raw, vno, code_raw, name_raw, unit_raw, qty_raw = row[:8]
        if code_raw is None and name_raw is None and qty_raw is None:
            continue
        c = norm_code(code_raw)
        q = sfloat(qty_raw)
        out_rows[r_idx] = (c, q)
        if c in cat:
            cat[c]["out"] += q

    wb.close()
    return cat, in_rows, out_rows

cat1, in1, out1 = load_state(P_BEFORE)
cat2, in2, out2 = load_state(P_AFTER)

# Build code mapping from transactions + normalized name matching
old_to_new = {}
for r in in1:
    if r in in2:
        c1, _ = in1[r]
        c2, _ = in2[r]
        if c1 and c2 and c1 != "2345":
            old_to_new[c1] = c2
for r in out1:
    if r in out2:
        c1, _ = out1[r]
        c2, _ = out2[r]
        if c1 and c2 and c1 != "2345":
            old_to_new[c1] = c2

# For items with 0 transactions, match by normalized name
name2_to_code = {" ".join(v["name"].split()): c for c, v in cat2.items()}
for c1, v1 in cat1.items():
    if c1 not in old_to_new:
        n_norm = " ".join(v1["name"].split())
        if n_norm in name2_to_code:
            old_to_new[c1] = name2_to_code[n_norm]

# Aggregate cat1 by target new code!
agg1 = defaultdict(lambda: {"old_codes": [], "old_names": [], "open": 0.0, "in": 0.0, "out": 0.0})
unmapped_old = []
for c1, v1 in cat1.items():
    c2 = old_to_new.get(c1)
    if not c2 or c2 not in cat2:
        unmapped_old.append((c1, v1))
    else:
        agg1[c2]["old_codes"].append(c1)
        agg1[c2]["old_names"].append(v1["name"])
        agg1[c2]["open"] += v1["open"]
        agg1[c2]["in"] += v1["in"]
        agg1[c2]["out"] += v1["out"]

lines = []
lines.append(f"Total active items BEFORE: {len(cat1)}, AFTER: {len(cat2)}")
lines.append(f"Unmapped old items: {len(unmapped_old)}")
for c1, v1 in unmapped_old:
    lines.append(f"  UNMAPPED OLD: [{c1}] '{v1['name']}' (open={v1['open']}, in={v1['in']}, out={v1['out']})")

diffs = []
merges = []
for c2, v2 in sorted(cat2.items()):
    a1 = agg1.get(c2, {"old_codes": [], "old_names": [], "open": 0.0, "in": 0.0, "out": 0.0})
    if len(a1["old_codes"]) > 1:
        merges.append((c2, v2["name"], list(zip(a1["old_codes"], a1["old_names"]))))
    b1 = a1["open"] + a1["in"] - a1["out"]
    b2 = v2["open"] + v2["in"] - v2["out"]
    if abs(a1["open"] - v2["open"]) > 1e-5 or abs(a1["in"] - v2["in"]) > 1e-5 or abs(a1["out"] - v2["out"]) > 1e-5:
        diffs.append((c2, v2["name"], a1, v2, b1, b2))

lines.append(f"\nMerged duplicate items ({len(merges)}):")
for c2, n2, old_list in merges:
    lines.append(f"  -> [{c2}] '{n2}' merged from: {old_list}")

lines.append(f"\nItems with ANY difference in (Open, In, Out, Bal) between BEFORE and AFTER ({len(diffs)}):")
for c2, n2, a1, v2, b1, b2 in diffs:
    lines.append(f"  [{c2}] '{n2}' (was {a1['old_codes']} {a1['old_names']}):")
    lines.append(f"      BEFORE: Open={a1['open']}, In={a1['in']}, Out={a1['out']}, Bal={b1}")
    lines.append(f"      AFTER : Open={v2['open']}, In={v2['in']}, Out={v2['out']}, Bal={b2}")

with open(r"c:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.1\wms-monorepo\scripts\scratch\step5_balance_check.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("Done! Lines:", len(lines))
