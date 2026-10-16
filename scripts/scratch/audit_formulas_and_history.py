# -*- coding: utf-8 -*-
import openpyxl
import json
import os
import sys

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'D:\The_Calm_WMS_V2.1\Excel_Database\The_Calm_WMS_V2.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=False)
ws = wb['الصرف']

print(f"Total rows in الصرف: {ws.max_row}")

no_formula_name = []
no_formula_unit = []

for r in range(4, ws.max_row + 1):
    val_e = ws.cell(row=r, column=5).value # كود الصنف
    if val_e is None:
        continue
    f_val = str(ws.cell(row=r, column=6).value or '')
    g_val = str(ws.cell(row=r, column=7).value or '')
    
    voucher = ws.cell(row=r, column=4).value # رقم الإذن
    date_val = ws.cell(row=r, column=2).value
    
    if not f_val.startswith('='):
        no_formula_name.append((r, voucher, val_e, f_val))
    if not g_val.startswith('='):
        no_formula_unit.append((r, voucher, val_e, g_val))

print(f"\nRows without formula in column F (اسم الصنف): {len(no_formula_name)}")
for item in no_formula_name:
    print(f"  Row {item[0]} | Voucher: {item[1]} | Code: {item[2]} | Val: {item[3]}")

print(f"\nRows without formula in column G (الوحدة): {len(no_formula_unit)}")
for item in no_formula_unit:
    print(f"  Row {item[0]} | Voucher: {item[1]} | Code: {item[2]} | Val: {item[3]}")

# Let's inspect formulas for rows around 2490 to 2646
sample_formulas = []
for r in range(2495, ws.max_row + 1):
    f_val = ws.cell(row=r, column=6).value
    g_val = ws.cell(row=r, column=7).value
    voucher = ws.cell(row=r, column=4).value
    sample_formulas.append((r, voucher, f_val, g_val))

print("\n--- Detailed formula inspect for recent dictated rows (2495..2646): ---")
for r, v, f, g in sample_formulas:
    f_is_form = str(f).startswith('=')
    g_is_form = str(g).startswith('=')
    if not f_is_form or not g_is_form or r in [2499, 2502, 2608, 2646]:
        print(f"Row {r} | V:{v} | F_form: {f_is_form} ({str(f)[:35]}) | G_form: {g_is_form} ({str(g)[:35]})")
