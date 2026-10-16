import sys
import os
import json
import time
import datetime
import urllib.request
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

EXCEL_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"
API_KEY = "AIzaSyBov0NrvjGouy7iByl4qGb6CpRVWsEc9CM"
PROJECT_ID = "the-calm-wms"
EMAIL = "mohamedalghala@gmail.com"
PASSWORD = "password123"

def make_request_with_retry(url, data=None, headers=None, retries=8, timeout=45):
    last_err = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, data=data, headers=headers or {})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            last_err = e
            print(f"      [Retry {attempt+1}/{retries}] Network warning: {e}. Backing off {2.0 * (attempt + 1):.1f}s...", flush=True)
            time.sleep(2.0 * (attempt + 1))
    raise last_err

def get_auth_token():
    auth_url = f"https://identitytoolkit.googleapis.com/v1/accounts:signInWithPassword?key={API_KEY}"
    payload = json.dumps({"email": EMAIL, "password": PASSWORD, "returnSecureToken": True}).encode("utf-8")
    headers = {"Content-Type": "application/json", "Connection": "close"}
    res = make_request_with_retry(auth_url, data=payload, headers=headers)
    return res["idToken"]

def commit_batch(token, writes):
    commit_url = f"https://firestore.googleapis.com/v1/projects/{PROJECT_ID}/databases/(default)/documents:commit"
    body = {"writes": writes}
    data = json.dumps(body).encode("utf-8")
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json", "Connection": "close"}
    return make_request_with_retry(commit_url, data=data, headers=headers)

def parse_excel_date(dt):
    if isinstance(dt, datetime.datetime):
        d_str = dt.strftime("%Y-%m-%d")
        ts_ms = int(dt.replace(tzinfo=datetime.timezone.utc).timestamp() * 1000)
        return d_str, ts_ms
    elif isinstance(dt, datetime.date):
        d_str = dt.strftime("%Y-%m-%d")
        dt_obj = datetime.datetime(dt.year, dt.month, dt.day, tzinfo=datetime.timezone.utc)
        return d_str, int(dt_obj.timestamp() * 1000)
    elif dt:
        s = str(dt).strip().split(" ")[0]
        try:
            if "-" in s:
                parts = [int(x) for x in s.split("-")]
                dt_obj = datetime.datetime(parts[0], parts[1], parts[2], tzinfo=datetime.timezone.utc)
                return dt_obj.strftime("%Y-%m-%d"), int(dt_obj.timestamp() * 1000)
            elif "/" in s:
                parts = [int(x) for x in s.split("/")]
                if parts[0] > 1000:
                    y, m, d = parts[0], parts[1], parts[2]
                else:
                    d, m, y = parts[0], parts[1], parts[2]
                    if y < 100: y += 2000
                dt_obj = datetime.datetime(y, m, d, tzinfo=datetime.timezone.utc)
                return dt_obj.strftime("%Y-%m-%d"), int(dt_obj.timestamp() * 1000)
        except Exception:
            pass
    return "2026-01-01", 1767225600000

def num_field(val):
    if float(val).is_integer():
        return {"integerValue": str(int(round(val)))}
    return {"doubleValue": round(float(val), 4)}

def run_sync_disbursements():
    print("=================================================================", flush=True)
    print("RESUMING FIRESTORE SYNC: DISBURSEMENTS COLLECTION", flush=True)
    print("=================================================================", flush=True)

    token = get_auth_token()
    print("1. Authenticated successfully with Firebase Cloud (`the-calm-wms`).", flush=True)

    print(f"\n2. Reading Excel catalog & disbursements from `{EXCEL_FILE}`...", flush=True)
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    
    ws_cat = wb["أرصدة المخزن"]
    catalog_map = {}
    for r in ws_cat.iter_rows(min_row=4, max_row=9405, values_only=True):
        c_code, c_dept, c_name, c_unit = r[:4]
        if c_name and str(c_name).strip():
            code_str = str(int(c_code) if isinstance(c_code, (int, float)) else str(c_code).strip())
            catalog_map[code_str] = {
                "name": str(c_name).strip(),
                "unit": str(c_unit or "عدد").strip(),
                "category": str(c_dept or "عام").strip()
            }

    ws_dis = wb["الصرف"]
    disbursements_list = []
    sum_out = 0.0
    for r in ws_dis.iter_rows(min_row=7, max_row=ws_dis.max_row, values_only=True):
        seq, dt, dept, inv, code, name, unit, qty, op, notes = r[:10]
        if seq is None:
            continue
        try:
            seq_int = int(seq)
        except Exception:
            continue
        code_str = str(int(code) if isinstance(code, (int, float)) else str(code or "").strip())
        qty_val = float(qty or 0)
        if not code_str or code_str == "9999" or qty_val <= 0:
            continue

        d_str, ts_ms = parse_excel_date(dt)
        inv_str = str(int(inv) if isinstance(inv, (int, float)) else str(inv or "").strip())
        dept_str = str(dept or "الصيانة").strip()
        cat_info = catalog_map.get(code_str, {})
        item_name = cat_info.get("name") or str(name or "").strip()
        item_unit = cat_info.get("unit") or str(unit or "عدد").strip()
        item_cat = cat_info.get("category") or "عام"
        notes_str = str(notes or "").strip()
        dest_str = notes_str if notes_str else "—"

        disb_obj = {
            "id": f"DISB_{seq_int}",
            "seq": seq_int,
            "date": d_str,
            "createdAt": ts_ms,
            "invoiceNo": inv_str,
            "department": dept_str,
            "destination": dest_str,
            "receiver": "—",
            "itemCode": code_str,
            "itemName": item_name,
            "category": item_cat,
            "unit": item_unit,
            "qty": qty_val,
            "notes": notes_str
        }
        disbursements_list.append(disb_obj)
        sum_out += qty_val
    wb.close()

    print(f"   -> Valid Disbursement Movements: {len(disbursements_list)}, Total Disbursed Qty: {sum_out:,.2f}", flush=True)

    print(f"\n3. Committing to Firestore `disbursements` collection in chunks of 200...", flush=True)
    CHUNK = 200
    dis_writes = []
    for d in disbursements_list:
        doc_path = f"projects/{PROJECT_ID}/databases/(default)/documents/disbursements/{d['id']}"
        dis_writes.append({
            "update": {
                "name": doc_path,
                "fields": {
                    "id": {"stringValue": d["id"]},
                    "seq": {"integerValue": str(d["seq"])},
                    "date": {"stringValue": d["date"]},
                    "createdAt": {"integerValue": str(d["createdAt"])},
                    "invoiceNo": {"stringValue": d["invoiceNo"]},
                    "department": {"stringValue": d["department"]},
                    "destination": {"stringValue": d["destination"]},
                    "receiver": {"stringValue": d["receiver"]},
                    "itemCode": {"stringValue": d["itemCode"]},
                    "itemName": {"stringValue": d["itemName"]},
                    "category": {"stringValue": d["category"]},
                    "unit": {"stringValue": d["unit"]},
                    "qty": num_field(d["qty"]),
                    "notes": {"stringValue": d["notes"]}
                }
            }
        })

    total_chunks = (len(dis_writes) + CHUNK - 1) // CHUNK
    for idx, i in enumerate(range(0, len(dis_writes), CHUNK)):
        chunk = dis_writes[i:i + CHUNK]
        res = commit_batch(token, chunk)
        committed_count = len(res.get('writeResults', []))
        print(f"   -> Batch {idx + 1}/{total_chunks}: {committed_count} documents committed.", flush=True)
        time.sleep(0.4)

    print("\n=================================================================", flush=True)
    print(f"SUCCESS: All {len(disbursements_list)} disbursement records synced to Firestore!", flush=True)
    print("=================================================================", flush=True)

if __name__ == "__main__":
    run_sync_disbursements()
