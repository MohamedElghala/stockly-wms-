import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

old_file = r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2_BACKUP_BEFORE_TREE_AND_8DIGIT_EXECUTION.xlsx'
new_file = r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx'

print("=== DEEP FORENSIC COMPARISON ===")
print("Old file:", old_file)
print("New file:", new_file)

wb_old = openpyxl.load_workbook(old_file, data_only=True)
wb_new = openpyxl.load_workbook(new_file, data_only=True)

# 1. Compare User's specific edits from yesterday:
print("\n--- 1. Checking User Edits from Yesterday ---")
ws_old_add = wb_old['الاضافة']
ws_new_add = wb_new['الاضافة']
ws_old_disb = wb_old['الصرف']
ws_new_disb = wb_new['الصرف']

# B1550 in الاضافة
print(f"Old الاضافة B1550: {ws_old_add['B1550'].value}")
print(f"New الاضافة B1550: {ws_new_add['B1550'].value}")

# C4 in الاضافة and الصرف
print(f"Old الاضافة C4: {repr(ws_old_add['C4'].value)} | New الاضافة C4: {repr(ws_new_add['C4'].value)}")
print(f"Old الصرف C4: {repr(ws_old_disb['C4'].value)} | New الصرف C4: {repr(ws_new_disb['C4'].value)}")

# Lotus remarks in الاضافة
lotus_old_count = sum(1 for r in range(7, ws_old_add.max_row + 1) if 'لوتس' in str(ws_old_add.cell(r, 9).value or ''))
lotus_new_count = sum(1 for r in range(7, ws_new_add.max_row + 1) if 'لوتس' in str(ws_new_add.cell(r, 9).value or ''))
print(f"Lotus remarks in الاضافة: Old = {lotus_old_count}, New = {lotus_new_count}")

# 2. Compare transaction counts and values in الاضافة
print("\n--- 2. Comparing Transactions in الاضافة ---")
old_add_rows = [r for r in range(7, ws_old_add.max_row + 1) if ws_old_add.cell(r, 1).value is not None]
new_add_rows = [r for r in range(7, ws_new_add.max_row + 1) if ws_new_add.cell(r, 1).value is not None]
print(f"Total numbered rows in الاضافة: Old = {len(old_add_rows)}, New = {len(new_add_rows)}")

old_add_qty = sum(float(ws_old_add.cell(r, 7).value or 0) for r in old_add_rows)
new_add_qty = sum(float(ws_new_add.cell(r, 7).value or 0) for r in new_add_rows)
print(f"Total Added Quantity: Old = {old_add_qty:,.2f}, New = {new_add_qty:,.2f} | Diff = {new_add_qty - old_add_qty:.4f}")

# 3. Compare transaction counts and values in الصرف
print("\n--- 3. Comparing Transactions in الصرف ---")
old_disb_rows = [r for r in range(7, ws_old_disb.max_row + 1) if ws_old_disb.cell(r, 1).value is not None]
new_disb_rows = [r for r in range(7, ws_new_disb.max_row + 1) if ws_new_disb.cell(r, 1).value is not None]
print(f"Total numbered rows in الصرف: Old = {len(old_disb_rows)}, New = {len(new_disb_rows)}")

old_disb_qty = sum(float(ws_old_disb.cell(r, 8).value or 0) for r in old_disb_rows)
new_disb_qty = sum(float(ws_new_disb.cell(r, 8).value or 0) for r in new_disb_rows)
print(f"Total Disbursed Quantity: Old = {old_disb_qty:,.2f}, New = {new_disb_qty:,.2f} | Diff = {new_disb_qty - old_disb_qty:.4f}")

# 4. Check Stock Balance Parity
print("\n--- 4. Stock Balance Parity ---")
ws_old_stock = wb_old['أرصدة المخزن']
ws_new_stock = wb_new['أرصدة المخزن']

old_open = sum(float(ws_old_stock.cell(r, 5).value or 0) for r in range(4, 8504))
new_open = sum(float(ws_new_stock.cell(r, 5).value or 0) for r in range(4, 9406))
print(f"Total Opening Balance: Old = {old_open:,.2f}, New = {new_open:,.2f} | Diff = {new_open - old_open:.4f}")

old_stock_tot = sum(float(ws_old_stock.cell(r, 8).value or 0) for r in range(4, 8504))
new_stock_tot = sum(float(ws_new_stock.cell(r, 8).value or 0) for r in range(4, 9406))
print(f"Total Current Stock: Old = {old_stock_tot:,.2f}, New = {new_stock_tot:,.2f} | Diff = {new_stock_tot - old_stock_tot:.4f}")
