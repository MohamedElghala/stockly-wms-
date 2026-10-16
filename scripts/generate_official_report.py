import sys
import os
import argparse
import datetime
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import win32com.client

sys.path.append(os.path.dirname(__file__))
from official_document_template import (
    create_official_document, add_official_signatures_footer,
    set_cell_background, set_cell_margins, set_table_borders,
    make_table_rtl, set_p_rtl, add_arabic_run
)

try:
    from config import EXCEL_FILE, REPORTS_DIR
except ImportError:
    EXCEL_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"
    REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "reports")
    os.makedirs(REPORTS_DIR, exist_ok=True)

def export_docx_to_pdf(docx_path, pdf_path):
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 is wdFormatPDF
        doc.Close()
        word.Quit()
        return True
    except Exception as e:
        print(f"   [Note] PDF export via Word COM: {e}")
        return False

def generate_voucher_report(voucher_num, voucher_type="add"):
    print(f"\n1. Fetching data for Voucher #{voucher_num}...")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    items = []
    v_date = "2026-09-30"
    v_entity = "موقع ستون رزيدنس"
    v_type_str = "إضافة وتوريد مخزني" if voucher_type == "add" else "صرف مخزني"

    if voucher_type == "add":
        ws = wb['الاضافة']
        for r in range(7, ws.max_row + 1):
            v = ws.cell(row=r, column=3).value
            if v == voucher_num or str(v).strip() == str(voucher_num):
                dt = ws.cell(row=r, column=2).value
                if dt:
                    v_date = dt.strftime("%Y-%m-%d") if isinstance(dt, datetime.datetime) else str(dt)[:10]
                code = ws.cell(row=r, column=4).value
                name = ws.cell(row=r, column=5).value
                unit = ws.cell(row=r, column=6).value or "-"
                qty = ws.cell(row=r, column=7).value or 0
                notes = ws.cell(row=r, column=9).value or "-"
                items.append({
                    "seq": len(items) + 1,
                    "code": code,
                    "name": name,
                    "unit": unit,
                    "qty": qty,
                    "notes": notes
                })
    else:
        ws = wb['الصرف']
        for r in range(7, ws.max_row + 1):
            v = ws.cell(row=r, column=4).value
            if v == voucher_num or str(v).strip() == str(voucher_num):
                dt = ws.cell(row=r, column=2).value
                if dt:
                    v_date = dt.strftime("%Y-%m-%d") if isinstance(dt, datetime.datetime) else str(dt)[:10]
                dept = ws.cell(row=r, column=3).value
                if dept:
                    v_entity = dept
                code = ws.cell(row=r, column=5).value
                name = ws.cell(row=r, column=6).value
                unit = ws.cell(row=r, column=7).value or "-"
                qty = ws.cell(row=r, column=8).value or 0
                notes = ws.cell(row=r, column=10).value or "-"
                items.append({
                    "seq": len(items) + 1,
                    "code": code,
                    "name": name,
                    "unit": unit,
                    "qty": qty,
                    "notes": notes
                })

    if not items:
        print(f"   [Error] No items found for voucher #{voucher_num} in sheet '{ws.title}'!")
        return None

    print(f"   Found {len(items)} items for voucher #{voucher_num}.")

    # Inspector assignment based on first item code or default
    first_code = str(items[0]["code"])
    if first_code.startswith("100"):
        inspector_role = "مشرف الكهرباء"
        inspector_name = "م / محمود متولي"
    elif first_code.startswith("101"):
        inspector_role = "مشرف السباكة"
        inspector_name = "م / حسام محمد"
    elif first_code.startswith("102"):
        inspector_role = "مشرف شبكات الري"
        inspector_name = "م / محمد حسن"
    else:
        inspector_role = "الفاحص والمعاين الفني"
        inspector_name = "م / محمود متولي"

    # Create Document
    doc, sig_info = create_official_document(
        subject=f"محضر فحص فني ومطابقة - إذن رقم {voucher_num}",
        date_str=v_date,
        inspector_role=inspector_role,
        inspector_name=inspector_name,
        warehouse_keeper_name="أ / محمد حلمي",
        project_manager_name="م / محمد عدوي"
    )

    # Document Title
    p_title = doc.add_paragraph()
    set_p_rtl(p_title)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    add_arabic_run(p_title, f"محضر فحص فني ومطابقة للمواصفات وتوريد مخزني", font_size=14, bold=True, color=RGBColor(31, 78, 121))

    p_sub = doc.add_paragraph()
    set_p_rtl(p_sub)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    add_arabic_run(p_sub, f"نوع الإذن: {v_type_str} | رقم الإذن: [ {voucher_num} ] | التاريخ: {v_date} | إجمالي البنود: {len(items)} صنفاً", font_size=10, italic=True, color=RGBColor(71, 85, 105))

    # Meta Table
    meta_tbl = doc.add_table(rows=2, cols=2)
    make_table_rtl(meta_tbl)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_tbl, color="CBD5E1", sz="4")

    meta_data = [
        [("المشروع:", "شركة ذا كالم - موقع ستون رزيدنس"), ("الجهة / المورد:", str(v_entity))],
        [("تاريخ الفحص والمطابقة:", v_date), (f"رقم إذن {v_type_str}:", f"إذن رقم {voucher_num} (عدد {len(items)} صنفاً)")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (lbl, val) in enumerate(row):
            c = meta_tbl.cell(r_idx, c_idx)
            set_cell_background(c, "F8FAFC" if c_idx == 0 else "FFFFFF")
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            p = c.paragraphs[0]
            set_p_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_arabic_run(p, lbl + " ", font_size=9, bold=True, color=RGBColor(31, 78, 121))
            add_arabic_run(p, val, font_size=9, bold=False, color=RGBColor(30, 41, 59))

    # Statement
    p_stmt = doc.add_paragraph()
    set_p_rtl(p_stmt)
    p_stmt.paragraph_format.space_before = Pt(8)
    p_stmt.paragraph_format.space_after = Pt(4)
    add_arabic_run(p_stmt, "بيان الأصناف المشمولة بالفحص الفني والمطابقة المخزنية:", font_size=10.5, bold=True, color=RGBColor(15, 23, 42))

    # Items Table
    tbl = doc.add_table(rows=len(items) + 1, cols=6)
    make_table_rtl(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4")

    headers = ["م", "كود الصنف", "اسم الصنف وتوصيفه", "الوحدة", "الكمية", "الملاحظات الفنية"]
    col_widths = [Inches(0.4), Inches(0.9), Inches(2.8), Inches(0.6), Inches(0.7), Inches(1.8)]

    # Header Row
    for col_idx, (h_text, w) in enumerate(zip(headers, col_widths)):
        cell = tbl.cell(0, col_idx)
        cell.width = w
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=60, bottom=60, left=40, right=40)
        p = cell.paragraphs[0]
        set_p_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_arabic_run(p, h_text, font_size=9.5, bold=True, color=RGBColor(255, 255, 255))

    # Data Rows
    for r_idx, item in enumerate(items, start=1):
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        row_vals = [
            str(item["seq"]),
            str(item["code"]),
            str(item["name"]),
            str(item["unit"]),
            str(item["qty"]),
            str(item["notes"])
        ]
        for c_idx, (val, w) in enumerate(zip(row_vals, col_widths)):
            cell = tbl.cell(r_idx, c_idx)
            cell.width = w
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
            p = cell.paragraphs[0]
            set_p_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3, 4] else WD_ALIGN_PARAGRAPH.RIGHT
            bold_flag = True if c_idx in [1, 4] else False
            add_arabic_run(p, val, font_size=9, bold=bold_flag, color=RGBColor(30, 41, 59))

    # Conclusion & Signatures
    p_concl = doc.add_paragraph()
    set_p_rtl(p_concl)
    p_concl.paragraph_format.space_before = Pt(8)
    p_concl.paragraph_format.space_after = Pt(4)
    add_arabic_run(p_concl, "تمت المعاينة والفحص الفني وثبتت مطابقة الأصناف المذكورة أعلاه للمواصفات الفنية المطلوبة، وتم إدراجها بالدفاتر المخزنية وفقاً للأصول.", font_size=9, italic=True, color=RGBColor(71, 85, 105))

    add_official_signatures_footer(doc, sig_info)

    # Save docx & pdf
    prefix = "محضر_فحص_توريد" if voucher_type == "add" else "محضر_فحص_صرف"
    docx_file = os.path.join(REPORTS_DIR, f"{prefix}_اذن_{voucher_num}.docx")
    pdf_file = os.path.join(REPORTS_DIR, f"{prefix}_اذن_{voucher_num}.pdf")

    doc.save(docx_file)
    print(f"   -> Word Document generated: {docx_file}")
    pdf_ok = export_docx_to_pdf(docx_file, pdf_file)
    if pdf_ok:
        print(f"   -> PDF Document generated: {pdf_file}")

    return docx_file

def generate_low_stock_report(threshold=2):
    print(f"\n1. Scanning for Low Stock Items (Balance <= {threshold})...")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb['أرصدة المخزن']

    low_items_by_dept = {}
    total_low = 0

    curr_dept = "عام"
    for r in range(4, 8504):
        code = ws.cell(row=r, column=1).value
        dept = ws.cell(row=r, column=2).value
        if dept:
            curr_dept = str(dept).strip()
        name = ws.cell(row=r, column=3).value
        unit = ws.cell(row=r, column=4).value or "-"
        bal = ws.cell(row=r, column=8).value

        if name is not None and str(name).strip() != "":
            try:
                bal_val = float(bal or 0)
            except:
                bal_val = 0.0

            if bal_val <= threshold:
                total_low += 1
                if curr_dept not in low_items_by_dept:
                    low_items_by_dept[curr_dept] = []
                low_items_by_dept[curr_dept].append({
                    "code": code,
                    "name": name,
                    "unit": unit,
                    "balance": bal_val
                })

    print(f"   Found {total_low} low-stock/critical items across {len(low_items_by_dept)} departments.")

    # Create Document
    doc, sig_info = create_official_document(
        subject="تقرير حصر النواقص وحدود إعادة الطلب للأصناف الحرجة",
        date_str=datetime.datetime.now().strftime("%Y-%m-%d"),
        inspector_role="أمين المخزن",
        inspector_name="أ / محمد حلمي",
        warehouse_keeper_name="أ / محمد حلمي",
        project_manager_name="م / محمد عدوي"
    )

    p_title = doc.add_paragraph()
    set_p_rtl(p_title)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    add_arabic_run(p_title, "تقرير حصر النواقص والأصناف الحرجة وحدود إعادة الطلب (Low Stock Alert)", font_size=13, bold=True, color=RGBColor(153, 27, 27))

    p_sub = doc.add_paragraph()
    set_p_rtl(p_sub)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    add_arabic_run(p_sub, f"المعايير: رصيد المخزن الحالي <= {threshold} | إجمالي الأصناف الحرجة: {total_low} صنفاً موزعة على {len(low_items_by_dept)} قسماً", font_size=10, italic=True, color=RGBColor(71, 85, 105))

    # Iterate departments
    for dept_name, items in low_items_by_dept.items():
        p_d = doc.add_paragraph()
        set_p_rtl(p_d)
        p_d.paragraph_format.space_before = Pt(6)
        p_d.paragraph_format.space_after = Pt(2)
        add_arabic_run(p_d, f"📌 قسم: {dept_name} (عدد {len(items)} صنفاً حرجاً)", font_size=10.5, bold=True, color=RGBColor(31, 78, 121))

        tbl = doc.add_table(rows=len(items) + 1, cols=5)
        make_table_rtl(tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="E2E8F0", sz="4")

        headers = ["م", "الكود", "اسم الصنف", "الوحدة", "الرصيد المتبقي"]
        col_w = [Inches(0.4), Inches(0.9), Inches(3.6), Inches(0.8), Inches(1.1)]

        for c_idx, (h, w) in enumerate(zip(headers, col_w)):
            cell = tbl.cell(0, c_idx)
            cell.width = w
            set_cell_background(cell, "7F1D1D")
            set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
            p = cell.paragraphs[0]
            set_p_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_arabic_run(p, h, font_size=9, bold=True, color=RGBColor(255, 255, 255))

        for r_idx, itm in enumerate(items, start=1):
            bg = "FEF2F2" if itm["balance"] <= 0 else "FFFFFF"
            vals = [str(r_idx), str(itm["code"]), str(itm["name"]), str(itm["unit"]), str(itm["balance"])]
            for c_idx, (v, w) in enumerate(zip(vals, col_w)):
                cell = tbl.cell(r_idx, c_idx)
                cell.width = w
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=30, bottom=30, left=30, right=30)
                p = cell.paragraphs[0]
                set_p_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3, 4] else WD_ALIGN_PARAGRAPH.RIGHT
                bold_flag = True if c_idx in [1, 4] else False
                font_c = RGBColor(185, 28, 28) if c_idx == 4 and itm["balance"] <= 0 else RGBColor(30, 41, 59)
                add_arabic_run(p, v, font_size=8.5, bold=bold_flag, color=font_c)

    # Signatures
    add_official_signatures_footer(doc, sig_info)

    docx_file = os.path.join(REPORTS_DIR, f"تقرير_حصر_نواقص_المخزن_{datetime.date.today()}.docx")
    pdf_file = os.path.join(REPORTS_DIR, f"تقرير_حصر_نواقص_المخزن_{datetime.date.today()}.pdf")
    doc.save(docx_file)
    print(f"   -> Word Document generated: {docx_file}")
    pdf_ok = export_docx_to_pdf(docx_file, pdf_file)
    if pdf_ok:
        print(f"   -> PDF Document generated: {pdf_file}")

    return docx_file

def main():
    parser = argparse.ArgumentParser(description="The Calm WMS - Official Document Generation Engine")
    parser.add_argument("--voucher", type=int, help="Voucher number to generate official report for")
    parser.add_argument("--type", choices=["add", "dis"], default="add", help="Voucher type: add (addition) or dis (disbursement)")
    parser.add_argument("--low-stock", action="store_true", help="Generate Low Stock / Reorder Alert report")
    parser.add_argument("--threshold", type=float, default=2.0, help="Reorder threshold for low stock (default <= 2)")

    args = parser.parse_args()

    print("==================================================")
    print("THE CALM WMS - OFFICIAL REPORTING ENGINE")
    print("==================================================")

    if args.voucher:
        generate_voucher_report(args.voucher, args.type)
    elif args.low_stock:
        generate_low_stock_report(args.threshold)
    else:
        # Interactive mode or demo test
        print("\nRunning demonstration tests:")
        print("Test 1: Generating Inspection Report for Addition Voucher #210...")
        generate_voucher_report(210, "add")
        print("\nTest 2: Generating Low Stock Report (Threshold <= 0)...")
        generate_low_stock_report(0)

    print("\n==================================================")
    print("REPORT ENGINE PROCESS COMPLETED!")
    print("==================================================")

if __name__ == "__main__":
    main()
