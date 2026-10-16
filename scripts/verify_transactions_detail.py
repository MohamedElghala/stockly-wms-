import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

old_file = r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2_BACKUP_BEFORE_TREE_AND_8DIGIT_EXECUTION.xlsx'
new_file = r'C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx'

with open(r'scripts\scratch\old_to_new_code_map.json', 'r', encoding='utf-8') as f:
    code_map = json.load(f)

wb_old = openpyxl.load_workbook(old_file, data_only=True)
wb_new = openpyxl.load_workbook(new_file, data_only=True)

# 1. Verify الاضافة rows 7 to 1726
ws_old_add = wb_old['الاضافة']
ws_new_add = wb_new['الاضافة']

add_mismatches = []
for r in range(7, 1727):
    # Col 1: Serial
    # Col 2: Date
    # Col 3: Permit
    # Col 4: Code
    # Col 5: Name
    # Col 6: Unit
    # Col 7: Qty
    # Col 8: Type
    # Col 9: Remarks
    old_c = str(ws_old_add.cell(r, 4).value or '').strip()
    new_c = str(ws_new_add.cell(r, 4).value or '').strip()
    
    expected_new = str(code_map.get(old_c, old_c))
    if new_c != expected_new:
        add_mismatches.append((r, 'Code', old_c, new_c, expected_new))
        
    old_qty = ws_old_add.cell(r, 7).value
    new_qty = ws_new_add.cell(r, 7).value
    if old_qty != new_qty:
        add_mismatches.append((r, 'Qty', old_qty, new_qty))
        
    old_date = str(ws_old_add.cell(r, 2).value or '')
    new_date = str(ws_new_add.cell(r, 2).value or '')
    if old_date != new_date:
        add_mismatches.append((r, 'Date', old_date, new_date))
        
    old_rem = str(ws_old_add.cell(r, 9).value or '')
    new_rem = str(ws_new_add.cell(r, 9).value or '')
    if old_rem != new_rem:
        add_mismatches.append((r, 'Remarks', old_rem, new_rem))

print(f"الاضافة transaction row-by-row checks: {len(add_mismatches)} mismatches.")
if add_mismatches:
    for m in add_mismatches[:5]:
        print("   Mismatch:", m)

# 2. Verify الصرف rows 7 to 2742
ws_old_disb = wb_old['الصرف']
ws_new_disb = wb_new['الصرف']

disb_mismatches = []
for r in range(7, 2743):
    # Col 1: Serial
    # Col 2: Date
    # Col 3: Dept/Recipient
    # Col 4: Permit
    # Col 5: Code
    # Col 6: Name
    # Col 7: Unit
    # Col 8: Qty
    # Col 9: Type
    # Col 10: Remarks
    old_c = str(ws_old_disb.cell(r, 5).value or '').strip()
    new_c = str(ws_new_disb.cell(r, 5).value or '').strip()
    
    expected_new = str(code_map.get(old_c, old_c))
    if new_c != expected_new:
        disb_mismatches.append((r, 'Code', old_c, new_c, expected_new))
        
    old_qty = ws_old_disb.cell(r, 8).value
    new_qty = ws_new_disb.cell(r, 8).value
    if old_qty != new_qty:
        disb_mismatches.append((r, 'Qty', old_qty, new_qty))
        
    old_date = str(ws_old_disb.cell(r, 2).value or '')
    new_date = str(ws_new_disb.cell(r, 2).value or '')
    if old_date != new_date:
        disb_mismatches.append((r, 'Date', old_date, new_date))
        
    old_dept = str(ws_old_disb.cell(r, 3).value or '')
    new_dept = str(ws_new_disb.cell(r, 3).value or '')
    if old_dept != new_dept:
        disb_mismatches.append((r, 'Dept', old_dept, new_dept))

print(f"الصرف transaction row-by-row checks: {len(disb_mismatches)} mismatches.")
if disb_mismatches:
    for m in disb_mismatches[:5]:
        print("   Mismatch:", m)

print("\n>>> All row-by-row transaction checks completed! <<<")
