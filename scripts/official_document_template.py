import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os, win32com.client

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_table_rtl(table):
    tblPr = table._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual')
    tblPr.append(bidi)

def set_p_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)

def add_arabic_run(p, text, font_size=11, bold=False, italic=False, color=RGBColor(0,0,0), font_name="Calibri"):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)
    return run

def create_official_document(
    subject="محضر فحص فني",
    date_str="28 / 9 / 2026",
    inspector_role="مشرف الكهرباء",
    inspector_name="م / محمود متولي",
    warehouse_keeper_name="أ / محمد حلمي",
    project_manager_name="م / محمد عدوي",
    output_docx_path=None
):
    doc = docx.Document()
    
    # Margins
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # 1. TOP HEADER TABLE (Two columns: Right = Metadata, Left = Logos)
    header_tbl = doc.add_table(rows=1, cols=2)
    make_table_rtl(header_tbl)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False

    # Right cell (Text metadata)
    cell_text = header_tbl.cell(0, 0)
    cell_text.width = Inches(4.8)
    p_hdr = cell_text.paragraphs[0]
    set_p_rtl(p_hdr)
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.paragraph_format.line_spacing = 1.3
    
    add_arabic_run(p_hdr, "شركة ذا كالم لإدارة المشروعات موقع ستون رزيدنس .\n", font_size=13, bold=True, color=RGBColor(20, 35, 60))
    add_arabic_run(p_hdr, f"الموضوع : {subject}\n", font_size=12, bold=True, color=RGBColor(31, 78, 121))
    add_arabic_run(p_hdr, f"التاريخ : {date_str}", font_size=11.5, bold=True, color=RGBColor(70, 80, 95))

    # Left cell (Company logos: The Calm & Stone Residence)
    cell_logo = header_tbl.cell(0, 1)
    cell_logo.width = Inches(2.2)
    p_logo = cell_logo.paragraphs[0]
    set_p_rtl(p_logo)
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    try:
        from config import LOGOS_DIR
        logos_dir = LOGOS_DIR
    except ImportError:
        logos_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "logos")
    
    calm_logo = os.path.join(logos_dir, "the_calm_logo.jpeg")
    stone_logo = os.path.join(logos_dir, "stone_residence_logo.jpeg")
    
    if os.path.exists(calm_logo):
        r_logo1 = p_logo.add_run()
        r_logo1.add_picture(calm_logo, width=Inches(1.3))
    p_logo2 = cell_logo.add_paragraph()
    set_p_rtl(p_logo2)
    p_logo2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(stone_logo):
        r_logo2 = p_logo2.add_run()
        r_logo2.add_picture(stone_logo, width=Inches(1.4))

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(8)
    pPr = p_div._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1F4E79"/></w:pBdr>')
    pPr.append(pBdr)

    return doc, {
        "inspector_role": inspector_role,
        "inspector_name": inspector_name,
        "warehouse_keeper_name": warehouse_keeper_name,
        "project_manager_name": project_manager_name
    }

def add_official_signatures_footer(doc, sig_info):
    p_sig_title = doc.add_paragraph()
    set_p_rtl(p_sig_title)
    p_sig_title.paragraph_format.space_before = Pt(14)
    p_sig_title.paragraph_format.space_after = Pt(4)

    sig_tbl = doc.add_table(rows=2, cols=3)
    make_table_rtl(sig_tbl)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    set_table_borders(sig_tbl, color="CBD5E1", sz="6")

    sig_columns = [
        ("المعاين والفاحص الفني", f"{sig_info['inspector_role']} : {sig_info['inspector_name']}"),
        ("أمين المخزن", f"امين المخازن : {sig_info['warehouse_keeper_name']}"),
        ("يعتمد / مدير المشروع", f"مدير المشروع : {sig_info['project_manager_name']}")
    ]

    for c_idx, (role, person) in enumerate(sig_columns):
        # Header cell
        cell_hdr = sig_tbl.cell(0, c_idx)
        set_cell_background(cell_hdr, "1F4E79")
        set_cell_margins(cell_hdr, top=60, bottom=60, left=60, right=60)
        p = cell_hdr.paragraphs[0]
        set_p_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_arabic_run(p, role, font_size=10.5, bold=True, color=RGBColor(255, 255, 255))

        # Body cell (Person name and Signature line)
        cell_bdy = sig_tbl.cell(1, c_idx)
        set_cell_background(cell_bdy, "FFFFFF")
        set_cell_margins(cell_bdy, top=100, bottom=100, left=60, right=60)
        p_b = cell_bdy.paragraphs[0]
        set_p_rtl(p_b)
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_arabic_run(p_b, f"{person}\n\n", font_size=9.5, bold=True, color=RGBColor(30, 41, 59))
        add_arabic_run(p_b, "التوقيع: ........................................", font_size=9, color=RGBColor(100, 116, 139))
