import sys
import os
import json
import time
import datetime
import urllib.request
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

try:
    from config import EXCEL_FILE, SCRIPT_DIR
except ImportError:
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    EXCEL_FILE = r"C:\Users\Mohamed Helmy\Downloads\The_Calm_WMS_V2.xlsx"

BACKUP_JSON = os.path.join(SCRIPT_DIR, "scratch", "live_firestore_backup_20261002.json")

API_KEY = "AIzaSyBov0NrvjGouy7iByl4qGb6CpRVWsEc9CM"
PROJECT_ID = "the-calm-wms"
EMAIL = "mohamedalghala@gmail.com"
PASSWORD = "password123"

DEPARTMENTS_17 = [
    {"name": "أدوات كهربائية", "order": 1, "codeRange": "100101 - 100319"},
    {"name": "أدوات سباكة", "order": 2, "codeRange": "101101 - 101383"},
    {"name": "أدوات خامات الري", "order": 3, "codeRange": "102101 - 102170"},
    {"name": "حمامات سباحة", "order": 4, "codeRange": "103101 - 103159"},
    {"name": "مستلزمات نظافة", "order": 5, "codeRange": "104101 - 104114"},
    {"name": "أدوات نقاشة", "order": 6, "codeRange": "105101 - 105140"},
    {"name": "أدوات وخامات نجارة", "order": 7, "codeRange": "106101 - 106151"},
    {"name": "أدوات عامة", "order": 8, "codeRange": "107101 - 107140"},
    {"name": "أدوات سيكوريت", "order": 9, "codeRange": "108102 - 108110"},
    {"name": "أدوات الوميتال", "order": 10, "codeRange": "109101 - 109120"},
    {"name": "أدوات مكاتب", "order": 11, "codeRange": "110101 - 110118"},
    {"name": "أدوات ميكانيكا", "order": 12, "codeRange": "111101 - 111109"},
    {"name": "أدوات الأمن وكاميرات المراقبة", "order": 13, "codeRange": "112101 - 112109"},
    {"name": "رخام ومواد بناء", "order": 14, "codeRange": "113101 - 113116"},
    {"name": "ملابس", "order": 15, "codeRange": "114101 - 114102"},
    {"name": "عدد و معدات", "order": 16, "codeRange": "115101 - 115135"},
    {"name": "قسم المستهلكات", "order": 17, "codeRange": "116101 - 116156"}
]

def get_auth_token():
    auth_url = f"https://identitytoolkit.googleapis.com/v1/accounts:signInWithPassword?key={API_KEY}"
    payload = json.dumps({"email": EMAIL, "password": PASSWORD, "returnSecureToken": True}).encode("utf-8")
    req = urllib.request.Request(auth_url, data=payload, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read().decode("utf-8"))
        return data["idToken"]

def commit_batch(token, writes):
    commit_url = f"https://firestore.googleapis.com/v1/projects/{PROJECT_ID}/databases/(default)/documents:commit"
    body = {"writes": writes}
    req = urllib.request.Request(
        commit_url,
        data=json.dumps(body).encode("utf-8"),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))

def sync_cloud():
    print("==================================================")
    print("STARTING PHASE 4: CLOUD SYNC & FIRESTORE ALIGNMENT")
    print("==================================================")

    # 1. Authenticate
    print("1. Authenticating with Firebase...")
    token = get_auth_token()
    print("   Authentication successful.")

    # 2. Read Master Excel Catalog
    print("\n2. Reading Master Excel Catalog (892 items)...")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb['أرصدة المخزن']

    catalog_items = []
    dept_counts = {}
    sum_bal = 0.0

    curr_dept = "عام"
    for r in range(4, 8504):
        c_code = ws.cell(row=r, column=1).value
        c_dept = ws.cell(row=r, column=2).value
        c_name = ws.cell(row=r, column=3).value
        c_unit = ws.cell(row=r, column=4).value or "عدد"
        c_init = float(ws.cell(row=r, column=5).value or 0)
        c_add = float(ws.cell(row=r, column=6).value or 0)
        c_dis = float(ws.cell(row=r, column=7).value or 0)
        c_bal = float(ws.cell(row=r, column=8).value or 0)

        if c_dept and str(c_dept).strip():
            curr_dept = str(c_dept).strip()

        if c_name and str(c_name).strip():
            code_str = str(int(c_code) if isinstance(c_code, (int, float)) else str(c_code).strip())
            name_str = str(c_name).strip()
            unit_str = str(c_unit).strip()

            catalog_items.append({
                "code": code_str,
                "name": name_str,
                "category": curr_dept,
                "unit": unit_str,
                "openingBalance": c_init,
                "totalIn": c_add,
                "totalOut": c_dis,
                "stock": c_bal
            })

            dept_counts[curr_dept] = dept_counts.get(curr_dept, 0) + 1
            sum_bal += c_bal

    print(f"   Read {len(catalog_items)} active catalog items.")
    print(f"   Total Balance: {sum_bal:,.2f}")
    assert len(catalog_items) == 892, f"Expected 892 items, got {len(catalog_items)}"
    assert abs(sum_bal - 15877.0) < 0.001, f"Expected 15,877.0, got {sum_bal}"

    # 3. Read live backup to identify obsolete items
    print("\n3. Comparing with existing Firestore documents...")
    with open(BACKUP_JSON, 'r', encoding='utf-8') as f:
        backup_data = json.load(f)

    existing_fs_items = backup_data['collections'].get('items', [])
    fs_code_set = set(it['_id'] for it in existing_fs_items)
    catalog_code_set = set(it["code"] for it in catalog_items)

    obsolete_codes = fs_code_set - catalog_code_set
    print(f"   Found {len(obsolete_codes)} obsolete items in Firestore to archive.")

    # 4. Sync 892 Active Catalog Items in Batches
    print("\n4. Uploading 892 Active Catalog Items to Firestore...")
    now_iso = datetime.datetime.now(datetime.timezone.utc).isoformat()
    now_ms = int(time.time() * 1000)

    writes_pool = []
    for item in catalog_items:
        doc_path = f"projects/{PROJECT_ID}/databases/(default)/documents/items/{item['code']}"
        status_val = "CRITICAL" if item["stock"] <= 0 else "ACTIVE"
        
        writes_pool.append({
            "update": {
                "name": doc_path,
                "fields": {
                    "id": {"stringValue": item["code"]},
                    "code": {"stringValue": item["code"]},
                    "name": {"stringValue": item["name"]},
                    "category": {"stringValue": item["category"]},
                    "unit": {"stringValue": item["unit"]},
                    "openingBalance": {"doubleValue": item["openingBalance"]},
                    "totalIn": {"doubleValue": item["totalIn"]},
                    "totalOut": {"doubleValue": item["totalOut"]},
                    "stock": {"doubleValue": item["stock"]},
                    "status": {"stringValue": status_val},
                    "classification": {"stringValue": status_val},
                    "alertLimit": {"integerValue": 10},
                    "updatedAt": {"integerValue": now_ms},
                    "syncedAt": {"stringValue": now_iso}
                }
            }
        })

    # Archive obsolete items
    for obs_code in obsolete_codes:
        doc_path = f"projects/{PROJECT_ID}/databases/(default)/documents/items/{obs_code}"
        writes_pool.append({
            "update": {
                "name": doc_path,
                "fields": {
                    "status": {"stringValue": "ARCHIVED"},
                    "classification": {"stringValue": "OBSOLETE"},
                    "stock": {"doubleValue": 0},
                    "updatedAt": {"integerValue": now_ms},
                    "syncedAt": {"stringValue": now_iso}
                }
            },
            "updateMask": {
                "fieldPaths": ["status", "classification", "stock", "updatedAt", "syncedAt"]
            }
        })

    # Execute in chunks of 200 writes
    CHUNK = 200
    for i in range(0, len(writes_pool), CHUNK):
        chunk = writes_pool[i:i + CHUNK]
        res = commit_batch(token, chunk)
        written = len(res.get("writeResults", []))
        print(f"   -> Batch {i//CHUNK + 1}/{(len(writes_pool)+CHUNK-1)//CHUNK}: Committed {written} writes.")

    print(f"   -> Successfully synchronized all items!")

    # 5. Populate 'categories' Collection with 17 Approved Departments
    print("\n5. Updating 'categories' Collection with the 17 Approved Departments...")
    cat_writes = []
    for d in DEPARTMENTS_17:
        cat_id = f"cat_{d['order']:02d}"
        doc_path = f"projects/{PROJECT_ID}/databases/(default)/documents/categories/{cat_id}"
        item_cnt = dept_counts.get(d["name"], 0)
        cat_writes.append({
            "update": {
                "name": doc_path,
                "fields": {
                    "id": {"stringValue": cat_id},
                    "name": {"stringValue": d["name"]},
                    "order": {"integerValue": d["order"]},
                    "codeRange": {"stringValue": d["codeRange"]},
                    "itemCount": {"integerValue": item_cnt},
                    "active": {"booleanValue": True},
                    "updatedAt": {"integerValue": now_ms}
                }
            }
        })
    cat_res = commit_batch(token, cat_writes)
    print(f"   -> Committed {len(cat_res.get('writeResults', []))} department categories.")

    # 6. Update Active Users' Departments in 'users'
    print("\n6. Updating Active Users' assigned departments...")
    all_dept_names = [d["name"] for d in DEPARTMENTS_17]
    user_writes = [
        # Abdelrahman Hassan (Manager) -> assign all 17 departments
        {
            "update": {
                "name": f"projects/{PROJECT_ID}/databases/(default)/documents/users/jjxzmm3E78UmZgpspIl1rOF9Ff92",
                "fields": {
                    "departments": {
                        "arrayValue": {
                            "values": [{"stringValue": name} for name in all_dept_names]
                        }
                    },
                    "department": {"stringValue": "كافة الأقسام"},
                    "category": {"stringValue": "كافة الأقسام"}
                }
            },
            "updateMask": {
                "fieldPaths": ["departments", "department", "category"]
            }
        },
        # Technician test 1 -> update plumbing & irrigation
        {
            "update": {
                "name": f"projects/{PROJECT_ID}/databases/(default)/documents/users/nQAd3sLYuvhtRGBZa56LeEbned22",
                "fields": {
                    "departments": {
                        "arrayValue": {
                            "values": [{"stringValue": "أدوات سباكة"}, {"stringValue": "أدوات خامات الري"}]
                        }
                    },
                    "department": {"stringValue": "أدوات سباكة"},
                    "category": {"stringValue": "أدوات سباكة"}
                }
            },
            "updateMask": {
                "fieldPaths": ["departments", "department", "category"]
            }
        }
    ]
    u_res = commit_batch(token, user_writes)
    print(f"   -> Updated {len(u_res.get('writeResults', []))} active users' departments.")

    print("\n==================================================")
    print("PHASE 4 CLOUD SYNC COMPLETED SUCCESSFULLY!")
    print("==================================================")

if __name__ == "__main__":
    sync_cloud()
